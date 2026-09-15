import os
import re
import time
import uuid
import sqlite3
import ipaddress
import subprocess
import concurrent.futures
from io import BytesIO
from datetime import datetime
from functools import wraps
from flask import Flask, render_template, request, redirect, url_for, flash, send_file, jsonify, session
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
import pandas as pd
from init_db import DB_PATH, EXCEL_PATH, NEW_EXCEL_PATH, BASE_DIR, init_db

app = Flask(__name__)
app.secret_key = 'bts_database_secret_key_antigravity'

UPLOAD_LOGS_DIR = os.path.join(BASE_DIR, 'data', 'uploaded_logs')
os.makedirs(UPLOAD_LOGS_DIR, exist_ok=True)

app.add_template_global(max, 'max')
app.add_template_global(min, 'min')



def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA cache_size = -65536;")  # 64 MB in-memory SQL cache
    conn.execute("PRAGMA temp_store = MEMORY;")   # Keep temporary tables and indices in RAM
    return conn


def login_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'user_id' not in session:
            flash('Please log in using your Username/Staff No and password to access the system.', 'warning')
            return redirect(url_for('login', next=request.url))
        return f(*args, **kwargs)
    return decorated_function


def row_to_dict(row):
    if row is None:
        return None
    d = dict(row)
    d['Location'] = d.get('location', '')
    d['cpan/maan/vsat'] = d.get('cpan_maan_vsat', '')
    d['tx-system-ip'] = d.get('tx_system_ip', '') or d.get('endpoint_ip', '') or d.get('cpan_a_end_ip', '')
    d['tx-system-location'] = d.get('tx_system_location', '') or d.get('endpoint_node_router', '') or d.get('cpan_a_end_node', '')
    
    raw_port = d.get('endpoint_ports', '') or d.get('cpan_a_end_ports', '') or d.get('tx_system_port', '')
    transformed_port = transform_tx_port(raw_port)
    
    if d.get('endpoint_ports'):
        d['endpoint_ports'] = transform_tx_port(d['endpoint_ports'])
    if d.get('cpan_a_end_ports'):
        d['cpan_a_end_ports'] = transform_tx_port(d['cpan_a_end_ports'])
    if d.get('tx_system_port'):
        d['tx_system_port'] = transformed_port
        
    d['tx-system-port'] = transformed_port
    d['vlan'] = d.get('vlan', '') or d.get('s1_c_vlan', '') or d.get('service_vlans', '') or d.get('oam_vlan', '')
    return d


def derive_ssa_and_location(site_name, existing_ssa="", existing_loc=""):
    if not site_name:
        return existing_ssa, existing_loc
    
    site_str = str(site_name).strip()
    if '_' in site_str:
        parts = site_str.split('_', 1)
        ssa_part = parts[0].strip()
        loc_part = parts[1].strip()
    else:
        ssa_part = site_str
        loc_part = ""
    
    m_full = re.match(r'^([A-Za-z0-9]*?\d+)(\s*[a-zA-Z].*)', ssa_part)
    if m_full:
        clean_ssa = m_full.group(1).strip()
        extra_loc = m_full.group(2).strip()
        if loc_part:
            loc_part = extra_loc + "_" + loc_part
        else:
            loc_part = extra_loc
        ssa_part = clean_ssa
        
    final_ssa = ssa_part if ssa_part else existing_ssa
    final_loc = loc_part if loc_part else existing_loc
    return final_ssa, final_loc


def transform_single_port(p_str):
    if not p_str:
        return ""
    p_str = str(p_str).strip()
    if not p_str or p_str.lower() in ('nan', 'none', '-'):
        return ""
        
    if "/" in p_str:
        tokens = [t.strip() for t in p_str.split("/") if t.strip()]
        if not tokens:
            return ""
            
        s_part = ""
        j_part = ""
        
        for t in tokens:
            if not s_part and re.match(r'^[Ss]\d+$', t):
                s_part = t.upper()
            elif not j_part and re.match(r'^[Jj]\d+$', t):
                j_part = t.upper()
                
        if not s_part and len(tokens) >= 1:
            s_part = tokens[0].upper()
            
        if not j_part:
            for t in tokens[1:]:
                if re.match(r'^[Pp]\d+$', t):
                    j_part = t.upper()
                    break
            if not j_part and len(tokens) >= 3:
                j_part = tokens[2].upper()
            elif not j_part and len(tokens) >= 2:
                j_part = tokens[1].upper()

        if s_part and j_part:
            return f"{s_part}/{j_part}"
        elif s_part:
            return s_part
        else:
            return p_str.upper()
    else:
        return p_str.strip().upper()


def transform_tx_port(port_str):
    if not port_str:
        return ""
    
    port_str = str(port_str).strip()
    if not port_str or port_str.lower() in ('nan', 'none', '-'):
        return ""
        
    delimiter = ";" if ";" in port_str else ("," if "," in port_str else None)
    
    if delimiter:
        parts = port_str.split(delimiter)
        transformed_parts = [transform_single_port(p) for p in parts]
        return (delimiter + " ").join(p for p in transformed_parts if p)
    else:
        return transform_single_port(port_str)


VALID_SEARCH_COLUMNS = {
    'all': 'Generic (All Columns)',
    'site_id': 'Site ID',
    'site_name': 'Site Name',
    'enodeb_address': 'eNodeB IP',
    'oam_cef_ip_pool': 'OAM CEF IP Pool',
    'ssa': 'SSA',
    'location': 'Location',
    'cpan_maan_vsat': 'Type (CPAN/MAAN/VSAT)',
    'mgmt_ip': 'Mgmt IP',
    'oam_vlan': 'OAM VLAN',
    'mgmt_rac_vlan': 'Mgmt / RAC VLAN',
    's1_c_vlan': 'S1-C VLAN',
    's1_u_vlan': 'S1-U VLAN',
    'endpoint_ip': 'Endpoint IP',
    'endpoint_node_router': 'Endpoint Node/Router',
    'cpan_a_end_node': 'CPAN A End Node',
    'cpan_a_end_ip': 'CPAN A End IP',
    'cpan_z_end_node': 'CPAN Z End Node',
    'cpan_z_end_ip': 'CPAN Z End IP',
    'tx_system_ip': 'TX System IP',
    'tx_system_location': 'TX System Location',
    'tx_system_port': 'TX System Port',
    'vlan': 'VLAN',
    'reason': 'Reason / Remark'
}


ALL_SEARCHABLE_COLS = [
    'site_id', 'site_name', 'enodeb_address', 'ssa', 'location', 'cpan_maan_vsat',
    'oam_vlan', 'mgmt_rac_vlan', 's1_c_vlan', 's1_u_vlan', 'mgmt_ip', 'mgmt_gateway',
    's1_u_ip', 'mme_ip', 'endpoint_type', 'endpoint_node_router', 'endpoint_ip',
    'l3_gateway_maan', 'endpoint_ports', 'cpan_a_end_node', 'cpan_a_end_ip',
    'cpan_a_end_ports', 'cpan_z_end_node', 'cpan_z_end_ip', 'cpan_service',
    'service_vlans', 'maan_l3_interface', 'maan_vpn', 'oam_cef_ip_pool',
    'oam_hw_gw', 'oam_hw_ip', 'tx_system_ip', 'tx_system_location', 'tx_system_port', 'vlan', 'reason'
]


VALID_SORT_COLUMNS = {
    'id': 'bts_sites.id',
    'site_id': 'site_id',
    'site_name': 'site_name',
    'ssa': 'ssa',
    'location': 'location',
    'cpan_maan_vsat': 'cpan_maan_vsat',
    'oam_cef_ip_pool': 'oam_cef_ip_pool',
    'enodeb_address': 'enodeb_address',
    'endpoint_ip': 'endpoint_ip',
    'endpoint_node_router': 'endpoint_node_router',
    'endpoint_ports': 'endpoint_ports',
    'oam_vlan': 'oam_vlan',
    'mgmt_rac_vlan': 'mgmt_rac_vlan',
    's1_c_vlan': 's1_c_vlan',
    's1_u_vlan': 's1_u_vlan',
    'mgmt_ip': 'mgmt_ip',
    'reason': 'reason',
    'created_at': 'created_at',
    'updated_at': 'updated_at'
}


def parse_search_query(query):
    if not query:
        return []
    raw = str(query).strip()
    if not raw:
        return []
    if any(c in raw for c in [',', ';', '\n', '\r', '\t']):
        items = [t.strip() for t in re.split(r'[,;\r\n\t]+', raw) if t.strip()]
    else:
        tokens = [t.strip() for t in raw.split() if t.strip()]
        if len(tokens) > 1 and all(re.match(r'^[A-Za-z0-9_\-\.\/:]+$', t) for t in tokens):
            items = tokens
        else:
            items = [raw]
    return list(dict.fromkeys(items))


app.add_template_global(parse_search_query, 'parse_multi_items')


def get_ssa_summary():
    conn = get_db()
    cursor = conn.cursor()
    sql = """
        SELECT 
            UPPER(TRIM(ssa)) as ssa_name,
            COUNT(*) as total_count,
            SUM(CASE WHEN LOWER(cpan_maan_vsat) LIKE '%cpan%' THEN 1 ELSE 0 END) as cpan_count,
            SUM(CASE WHEN LOWER(cpan_maan_vsat) LIKE '%maan%' THEN 1 ELSE 0 END) as maan_count,
            SUM(CASE WHEN LOWER(cpan_maan_vsat) NOT LIKE '%cpan%' AND LOWER(cpan_maan_vsat) NOT LIKE '%maan%' THEN 1 ELSE 0 END) as other_count
        FROM bts_sites 
        WHERE ssa IS NOT NULL AND TRIM(ssa) != '' 
        GROUP BY UPPER(TRIM(ssa))
        ORDER BY total_count DESC, ssa_name ASC;
    """
    cursor.execute(sql)
    rows = cursor.fetchall()
    conn.close()
    
    summary = []
    for r in rows:
        summary.append({
            'ssa_name': r['ssa_name'] or 'UNKNOWN',
            'total_count': r['total_count'] or 0,
            'cpan_count': r['cpan_count'] or 0,
            'maan_count': r['maan_count'] or 0,
            'other_count': r['other_count'] or 0
        })
    return summary


def search_db(query="", search_by="all", selected_ssa="", page=1, per_page=25, sort_by="id", sort_order="asc"):
    conn = get_db()
    cursor = conn.cursor()
    
    if search_by not in VALID_SEARCH_COLUMNS:
        search_by = "all"
        
    sort_col = VALID_SORT_COLUMNS.get(sort_by, 'bts_sites.id')
    sort_dir = "DESC" if str(sort_order).lower() in ("desc", "descending") else "ASC"
    
    cursor.execute("SELECT COUNT(*) FROM bts_sites;")
    total_records = cursor.fetchone()[0]
    
    cursor.execute("SELECT COUNT(*) FROM bts_sites WHERE LOWER(cpan_maan_vsat) LIKE '%cpan%';")
    cpan_total_count = cursor.fetchone()[0]
    
    cursor.execute("SELECT COUNT(*) FROM bts_sites WHERE LOWER(cpan_maan_vsat) LIKE '%maan%';")
    maan_total_count = cursor.fetchone()[0]
    
    cursor.execute("SELECT DISTINCT ssa FROM bts_sites WHERE ssa IS NOT NULL AND ssa != '' ORDER BY ssa ASC;")
    unique_ssas = [r['ssa'] for r in cursor.fetchall()]
    
    ssa_summary = get_ssa_summary()
    
    items = parse_search_query(query)
    
    if len(items) > 1:
        ssa_where = ""
        ssa_params = []
        if selected_ssa:
            ssa_where = "WHERE LOWER(ssa) = LOWER(?)"
            ssa_params = [selected_ssa.strip()]
            
        if search_by != "all" and search_by in ALL_SEARCHABLE_COLS:
            text_expr = f"IFNULL({search_by}, '')"
        else:
            text_expr = " || ' ' || ".join([f"IFNULL({c}, '')" for c in ALL_SEARCHABLE_COLS])
            
        sql = f"SELECT id, cpan_maan_vsat, ({text_expr}) as search_text FROM bts_sites {ssa_where} ORDER BY {sort_col} {sort_dir};"
        cursor.execute(sql, ssa_params)
        all_rows = cursor.fetchall()
        
        pattern = re.compile('|'.join([re.escape(it) for it in items]), re.IGNORECASE)
        
        matched_ids = []
        cpan_filtered_count = 0
        maan_filtered_count = 0
        
        for r in all_rows:
            if pattern.search(r['search_text']):
                matched_ids.append(r['id'])
                t_val = (r['cpan_maan_vsat'] or '').lower()
                if 'cpan' in t_val:
                    cpan_filtered_count += 1
                if 'maan' in t_val:
                    maan_filtered_count += 1
                    
        filtered_count = len(matched_ids)
        total_pages = max(1, (filtered_count + per_page - 1) // per_page)
        page = max(1, min(page, total_pages))
        offset = (page - 1) * per_page
        
        page_ids = matched_ids[offset : offset + per_page]
        
        if page_ids:
            placeholders = ','.join(['?']*len(page_ids))
            page_sql = f"SELECT * FROM bts_sites WHERE id IN ({placeholders}) ORDER BY {sort_col} {sort_dir};"
            cursor.execute(page_sql, page_ids)
            rows = cursor.fetchall()
        else:
            rows = []
    else:
        where_clauses = []
        params = []
        
        if query:
            q_like = f"%{query.strip()}%"
            if search_by != "all" and search_by in ALL_SEARCHABLE_COLS:
                where_clauses.append(f"LOWER({search_by}) LIKE LOWER(?)")
                params.append(q_like)
            else:
                or_clauses = [f"LOWER({c}) LIKE LOWER(?)" for c in ALL_SEARCHABLE_COLS]
                where_clauses.append("(" + " OR ".join(or_clauses) + ")")
                params.extend([q_like] * len(ALL_SEARCHABLE_COLS))
                
        if selected_ssa:
            where_clauses.append("LOWER(ssa) = LOWER(?)")
            params.append(selected_ssa.strip())
            
        where_sql = ""
        if where_clauses:
            where_sql = "WHERE " + " AND ".join(where_clauses)
            
        count_sql = f"SELECT COUNT(*) FROM bts_sites {where_sql};"
        cursor.execute(count_sql, params)
        filtered_count = cursor.fetchone()[0]
        
        cpan_where = list(where_clauses) + ["LOWER(cpan_maan_vsat) LIKE '%cpan%'"]
        cpan_where_sql = "WHERE " + " AND ".join(cpan_where)
        cursor.execute(f"SELECT COUNT(*) FROM bts_sites {cpan_where_sql};", params)
        cpan_filtered_count = cursor.fetchone()[0]
        
        maan_where = list(where_clauses) + ["LOWER(cpan_maan_vsat) LIKE '%maan%'"]
        maan_where_sql = "WHERE " + " AND ".join(maan_where)
        cursor.execute(f"SELECT COUNT(*) FROM bts_sites {maan_where_sql};", params)
        maan_filtered_count = cursor.fetchone()[0]
        
        total_pages = max(1, (filtered_count + per_page - 1) // per_page)
        page = max(1, min(page, total_pages))
        offset = (page - 1) * per_page
        
        data_sql = f"SELECT * FROM bts_sites {where_sql} ORDER BY {sort_col} {sort_dir} LIMIT ? OFFSET ?;"
        cursor.execute(data_sql, params + [per_page, offset])
        rows = cursor.fetchall()

    conn.close()
    
    records = [row_to_dict(r) for r in rows]
    return {
        'records': records,
        'query': query,
        'search_by': search_by,
        'selected_ssa': selected_ssa,
        'sort_by': sort_by,
        'sort_order': sort_order,
        'unique_ssas': unique_ssas,
        'ssa_summary': ssa_summary,
        'page': page,
        'per_page': per_page,
        'total_pages': total_pages,
        'total_records': total_records,
        'filtered_count': filtered_count,
        'cpan_total_count': cpan_total_count,
        'maan_total_count': maan_total_count,
        'cpan_filtered_count': cpan_filtered_count,
        'maan_filtered_count': maan_filtered_count
    }


# ==========================================
# AUTHENTICATION & USER REGISTRATION ROUTES
# ==========================================

@app.route('/login', methods=['GET', 'POST'])
def login():
    if session.get('user_id'):
        return redirect(url_for('index'))
        
    if request.method == 'POST':
        user_input = request.form.get('username', '').strip()
        password = request.form.get('password', '').strip()
        
        if not user_input or not password:
            flash('Please enter both Username/Staff Number and Password.', 'danger')
            return render_template('login.html')
            
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM users WHERE LOWER(username) = LOWER(?) OR LOWER(staff_no) = LOWER(?);", (user_input, user_input))
        user = cursor.fetchone()
        conn.close()
        
        if user and check_password_hash(user['password_hash'], password):
            session['user_id'] = user['id']
            session['username'] = user['username'] or user['staff_no']
            session['staff_no'] = user['staff_no']
            session['full_name'] = user['full_name']
            session['role'] = user['role']
            
            flash(f"Welcome back, {user['full_name']}!", "success")
            next_url = request.args.get('next')
            if next_url and next_url.startswith('/'):
                return redirect(next_url)
            return redirect(url_for('index'))
        else:
            flash('Invalid Username/Staff Number or Password. Please try again.', 'danger')
            
    return render_template('login.html')


@app.route('/register', methods=['GET', 'POST'])
def register():
    if session.get('user_id'):
        return redirect(url_for('index'))
        
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        staff_no = request.form.get('staff_no', '').strip()
        full_name = request.form.get('full_name', '').strip()
        email = request.form.get('email', '').strip()
        department = request.form.get('department', '').strip()
        password = request.form.get('password', '').strip()
        confirm_password = request.form.get('confirm_password', '').strip()
        
        if not username or not staff_no or not full_name or not password:
            flash('Username, Staff Number, Full Name, and Password are required.', 'danger')
            return render_template('register.html', form_data=request.form)
            
        if password != confirm_password:
            flash('Password and Confirm Password do not match.', 'danger')
            return render_template('register.html', form_data=request.form)
            
        conn = get_db()
        cursor = conn.cursor()
        
        cursor.execute("SELECT id FROM users WHERE LOWER(username) = LOWER(?) OR LOWER(staff_no) = LOWER(?);", (username, staff_no))
        if cursor.fetchone() is not None:
            conn.close()
            flash(f'Username "{username}" or Staff Number "{staff_no}" is already registered.', 'danger')
            return render_template('register.html', form_data=request.form)
            
        pass_hash = generate_password_hash(password)
        cursor.execute('''
            INSERT INTO users (username, staff_no, full_name, email, department, password_hash, role)
            VALUES (?, ?, ?, ?, ?, ?, ?);
        ''', (username, staff_no, full_name, email, department, pass_hash, 'staff'))
        
        conn.commit()
        conn.close()
        
        flash(f'Account for "{username}" created successfully! Please log in with your credentials.', 'success')
        return redirect(url_for('login'))
        
    return render_template('register.html', form_data={})


@app.route('/logout')
def logout():
    session.clear()
    flash('You have been logged out successfully.', 'info')
    return redirect(url_for('login'))


# ==========================================
# USER PROFILE ROUTES
# ==========================================

@app.route('/profile')
@login_required
def profile():
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM users WHERE id = ?;", (session['user_id'],))
    user_row = cursor.fetchone()
    conn.close()
    
    if not user_row:
        session.clear()
        flash('User account not found.', 'danger')
        return redirect(url_for('login'))
        
    return render_template('profile.html', user=dict(user_row))


@app.route('/profile/update', methods=['POST'])
@login_required
def update_profile():
    username = request.form.get('username', '').strip()
    full_name = request.form.get('full_name', '').strip()
    email = request.form.get('email', '').strip()
    department = request.form.get('department', '').strip()
    
    if not username or not full_name:
        flash('Username and Full Name are required.', 'danger')
        return redirect(url_for('profile'))
        
    conn = get_db()
    cursor = conn.cursor()
    
    cursor.execute("SELECT id FROM users WHERE LOWER(username) = LOWER(?) AND id != ?;", (username, session['user_id']))
    if cursor.fetchone() is not None:
        conn.close()
        flash(f'Username "{username}" is already taken by another user.', 'danger')
        return redirect(url_for('profile'))
        
    cursor.execute('''
        UPDATE users SET username = ?, full_name = ?, email = ?, department = ? WHERE id = ?;
    ''', (username, full_name, email, department, session['user_id']))
    conn.commit()
    conn.close()
    
    session['username'] = username
    session['full_name'] = full_name
    flash('Your profile details have been updated successfully!', 'success')
    return redirect(url_for('profile'))


@app.route('/profile/change-password', methods=['POST'])
@login_required
def change_password():
    current_password = request.form.get('current_password', '').strip()
    new_password = request.form.get('new_password', '').strip()
    confirm_password = request.form.get('confirm_password', '').strip()
    
    if not current_password or not new_password or not confirm_password:
        flash('All password fields are required.', 'danger')
        return redirect(url_for('profile'))
        
    if new_password != confirm_password:
        flash('New Password and Confirm Password do not match.', 'danger')
        return redirect(url_for('profile'))
        
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT password_hash FROM users WHERE id = ?;", (session['user_id'],))
    row = cursor.fetchone()
    
    if not row or not check_password_hash(row['password_hash'], current_password):
        conn.close()
        flash('Incorrect Current Password. Password was not changed.', 'danger')
        return redirect(url_for('profile'))
        
    new_hash = generate_password_hash(new_password)
    cursor.execute("UPDATE users SET password_hash = ? WHERE id = ?;", (new_hash, session['user_id']))
    conn.commit()
    conn.close()
    
    flash('Your password has been updated successfully!', 'success')
    return redirect(url_for('profile'))


# ==========================================
# SITE & DATABASE ROUTES (PROTECTED)
# ==========================================

@app.route('/')
@login_required
def index():
    query = request.args.get('q', '').strip()
    search_by = request.args.get('search_by', 'all').strip()
    selected_ssa = request.args.get('ssa', '').strip()
    sort_by = request.args.get('sort_by', 'id').strip()
    sort_order = request.args.get('sort_order', 'asc').strip()
    page = request.args.get('page', 1, type=int)
    per_page = request.args.get('per_page', 25, type=int)
    
    data = search_db(query=query, search_by=search_by, selected_ssa=selected_ssa, page=page, per_page=per_page, sort_by=sort_by, sort_order=sort_order)
    return render_template('index.html', **data)


@app.route('/api/search')
@login_required
def api_search():
    query = request.args.get('q', '').strip()
    search_by = request.args.get('search_by', 'all').strip()
    selected_ssa = request.args.get('ssa', '').strip()
    sort_by = request.args.get('sort_by', 'id').strip()
    sort_order = request.args.get('sort_order', 'asc').strip()
    page = request.args.get('page', 1, type=int)
    per_page = request.args.get('per_page', 25, type=int)
    
    data = search_db(query=query, search_by=search_by, selected_ssa=selected_ssa, page=page, per_page=per_page, sort_by=sort_by, sort_order=sort_order)
    return jsonify(data)


@app.route('/api/site/<path:site_id>')
@login_required
def api_get_site(site_id):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM bts_sites WHERE LOWER(site_id) = LOWER(?);", (site_id,))
    row = cursor.fetchone()
    conn.close()
    
    if not row:
        return jsonify({'success': False, 'message': f'Site ID "{site_id}" not found.'}), 404
        
    return jsonify({'success': True, 'site': row_to_dict(row)})


def log_activity(action, target_type="", target_id="", details=""):
    try:
        user_id = session.get('user_id')
        username = session.get('username') or session.get('full_name') or 'System'
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute('''
            INSERT INTO activity_logs (user_id, username, action, target_type, target_id, details)
            VALUES (?, ?, ?, ?, ?, ?);
        ''', (user_id, username, action, target_type, target_id, details))
        conn.commit()
        conn.close()
    except Exception as e:
        print(f"Error logging activity: {e}")


@app.route('/api/site/<site_id>/update-reason', methods=['POST'])
@login_required
def api_update_site_reason(site_id):
    data = request.get_json(silent=True) or request.form
    reason = str(data.get('reason', '')).strip()
    
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("UPDATE bts_sites SET reason = ?, updated_at = CURRENT_TIMESTAMP WHERE LOWER(site_id) = LOWER(?);", (reason, site_id))
    conn.commit()
    updated = cursor.rowcount > 0
    conn.close()
    
    if not updated:
        return jsonify({'success': False, 'message': f'Site ID "{site_id}" not found.'}), 404
        
    log_activity('UPDATE_REASON', 'bts_sites', site_id, f"Updated reason/remark to: {reason}")
    return jsonify({'success': True, 'site_id': site_id, 'reason': reason, 'message': 'Reason updated successfully.'})


@app.route('/api/sites/bulk-update-reason', methods=['POST'])
@login_required
def api_bulk_update_reason():
    data = request.get_json(silent=True) or request.form
    site_ids = data.get('site_ids', [])
    reason = str(data.get('reason', '')).strip()
    
    if isinstance(site_ids, str):
        site_ids = [s.strip() for s in site_ids.split(',') if s.strip()]
        
    if not isinstance(site_ids, list) or not site_ids:
        return jsonify({'success': False, 'message': 'No site IDs provided.'}), 400
        
    conn = get_db()
    cursor = conn.cursor()
    
    placeholders = ','.join(['?'] * len(site_ids))
    sql = f"UPDATE bts_sites SET reason = ?, updated_at = CURRENT_TIMESTAMP WHERE site_id IN ({placeholders});"
    cursor.execute(sql, [reason] + site_ids)
    conn.commit()
    updated_count = cursor.rowcount
    conn.close()
    
    log_activity(
        action='BULK_UPDATE_REASON',
        target_type='bts_sites',
        target_id=f"{updated_count} sites",
        details=f"Updated reason to '{reason}' for {updated_count} sites: {', '.join(site_ids[:10])}{'...' if len(site_ids)>10 else ''}"
    )
    
    return jsonify({
        'success': True,
        'updated_count': updated_count,
        'reason': reason,
        'message': f'Successfully updated reason for {updated_count} site(s).'
    })



@app.route('/api/activity-logs')
@login_required
def api_activity_logs():
    limit = request.args.get('limit', 25, type=int)
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM activity_logs ORDER BY created_at DESC LIMIT ?;", (limit,))
    logs = [dict(r) for r in cursor.fetchall()]
    conn.close()
    return jsonify({'success': True, 'logs': logs})
@app.route('/api/sites/summary')
@login_required
def api_sites_summary():
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT COUNT(*) FROM bts_sites;")
    total_sites = cursor.fetchone()[0]
    cursor.execute("SELECT COUNT(*) FROM bts_sites WHERE LOWER(cpan_maan_vsat) LIKE '%cpan%';")
    cpan_sites = cursor.fetchone()[0]
    cursor.execute("SELECT COUNT(*) FROM bts_sites WHERE LOWER(cpan_maan_vsat) LIKE '%maan%';")
    maan_sites = cursor.fetchone()[0]
    cursor.execute("SELECT COUNT(*) FROM cpan_nodes;")
    cpan_nodes_count = cursor.fetchone()[0]
    conn.close()
    
    ssa_summary = get_ssa_summary()
    return jsonify({
        'success': True,
        'total_sites': total_sites,
        'cpan_sites': cpan_sites,
        'maan_sites': maan_sites,
        'cpan_nodes_count': cpan_nodes_count,
        'ssa_counts': ssa_summary
    })



def check_ip_reachability(target_ip, returncode, output):
    if returncode != 0 or not output:
        return False

    output_upper = output.upper()
    clean_target = str(target_ip).split('/')[0].strip().upper()

    failure_keywords = [
        'UNREACHABLE', 'TIMED OUT', '100% LOSS', '100% PACKET LOSS',
        'GENERAL FAILURE', 'TRANSMIT FAILED', 'HARDWARE ERROR',
        'EXPIRED IN TRANSIT', 'UNKNOWN HOST', 'COULD NOT FIND HOST',
        'ADMINISTRATIVELY PROHIBITED'
    ]
    if any(kw in output_upper for kw in failure_keywords):
        return False

    has_echo_reply_marker = ('TTL=' in output_upper or 'BYTES=' in output_upper or 'BYTES FROM' in output_upper)
    if not has_echo_reply_marker:
        return False

    for line in output_upper.splitlines():
        if 'REPLY FROM' in line:
            if clean_target in line and ('TTL=' in line or 'BYTES=' in line or 'TIME' in line):
                return True
            if clean_target not in line:
                return False

    return has_echo_reply_marker


@app.route('/api/ping', methods=['POST'])
@login_required
def api_ping():
    data = request.get_json(silent=True) or request.form
    raw_ip = str(data.get('ip', '')).strip()
    
    if not raw_ip:
        return jsonify({'success': False, 'message': 'IP address is required.'}), 400
        
    clean_ip = raw_ip.split('/')[0].strip()
    
    try:
        ip_obj = ipaddress.ip_address(clean_ip)
    except ValueError:
        return jsonify({
            'success': False,
            'message': f'"{raw_ip}" is not a valid IPv4/IPv6 address.'
        }), 400
        
    if os.name == 'nt':
        cmd = ['ping', '-n', '4', '-w', '1000', str(ip_obj)]
    else:
        cmd = ['ping', '-c', '4', '-W', '1', str(ip_obj)]
    
    try:
        proc = subprocess.run(cmd, capture_output=True, text=True, timeout=10)
        output = (proc.stdout or proc.stderr or "No ping output returned.").strip()
        is_reachable = check_ip_reachability(str(ip_obj), proc.returncode, output)
        
        return jsonify({
            'success': True,
            'ip': str(ip_obj),
            'raw_ip': raw_ip,
            'output': output,
            'message': output,
            'is_reachable': is_reachable,
            'exit_code': proc.returncode
        })
    except subprocess.TimeoutExpired:
        timeout_msg = f"Ping request to {ip_obj} timed out (exceeded 10s timeout)."
        return jsonify({
            'success': True,
            'ip': str(ip_obj),
            'raw_ip': raw_ip,
            'output': timeout_msg,
            'message': timeout_msg,
            'is_reachable': False,
            'exit_code': -1
        })
    except Exception as e:
        err_msg = f"Error executing ping: {str(e)}"
        return jsonify({
            'success': False,
            'output': err_msg,
            'message': err_msg
        }), 500


# ==========================================
# LOG ANALYZER & MULTI-FORMAT SEARCH SYSTEM
# ==========================================

def read_log_file_lines(file_path):
    """
    Reads lines from a log file (supports .txt, .log, .csv, .json, .xlsx).
    Returns a list of string lines with 1-based line indexing.
    """
    if not os.path.exists(file_path):
        return []

    ext = os.path.splitext(file_path)[1].lower()
    lines = []

    try:
        if ext in ('.xlsx', '.xls'):
            df = pd.read_excel(file_path, dtype=str).fillna('')
            lines = [", ".join(f"{col}: {val}" for col, val in row.items() if str(val).strip()) for _, row in df.iterrows()]
        elif ext == '.json':
            import json as json_mod
            with open(file_path, 'r', encoding='utf-8', errors='replace') as f:
                content = f.read()
                try:
                    data = json_mod.loads(content)
                    if isinstance(data, list):
                        for item in data:
                            lines.append(json_mod.dumps(item))
                    elif isinstance(data, dict):
                        for k, v in data.items():
                            lines.append(f"{k}: {json_mod.dumps(v)}")
                    else:
                        lines = content.splitlines()
                except Exception:
                    lines = content.splitlines()
        else:
            with open(file_path, 'r', encoding='utf-8', errors='replace') as f:
                lines = [line.rstrip('\r\n') for line in f]
    except Exception as e:
        print(f"Error reading log file {file_path}: {e}")
        lines = [f"[File Read Error]: {str(e)}"]

    return lines


def extract_entities_from_matches(matches):
    """
    Extracts IP addresses and VLAN IDs from search matches.
    """
    ip_pattern = re.compile(r'\b(?:\d{1,3}\.){3}\d{1,3}(?:/\d{1,2})?\b')
    vlan_pattern = re.compile(r'\b(?:VLAN|vlan|Vlan)[\s:=_-]*(\d{1,5})\b|\bVLAN\s*(\d{1,5})\b', re.IGNORECASE)

    ips = set()
    vlans = set()

    for m in matches:
        text = m.get('line_text', '')
        for ip in ip_pattern.findall(text):
            clean_ip = ip.split('/')[0]
            parts = clean_ip.split('.')
            if len(parts) == 4 and all(p.isdigit() and 0 <= int(p) <= 255 for p in parts):
                ips.add(ip)

        for match_groups in vlan_pattern.findall(text):
            for v in match_groups:
                if v:
                    vlans.add(v)

    return {
        'ips': sorted(list(ips)),
        'vlans': sorted(list(vlans), key=lambda x: int(x) if x.isdigit() else x)
    }


TIMESTAMP_REGEX_PATTERNS = [
    re.compile(r'\b(\d{4}-\d{2}-\d{2}[ T]\d{2}:\d{2}(?::\d{2})?)\b'),
    re.compile(r'\b(\d{2}[/-]\d{2}[/-]\d{4}[ T]\d{2}:\d{2}(?::\d{2})?)\b'),
    re.compile(r'\b(\d{4}/\d{2}/\d{2}[ T]\d{2}:\d{2}(?::\d{2})?)\b'),
    re.compile(r'\b([A-Z][a-z]{2}\s+\d{1,2}\s+\d{2}:\d{2}:\d{2})\b')
]

def parse_date_str(date_str):
    if not date_str:
        return None
    date_str = str(date_str).strip().replace('T', ' ')
    for fmt in ('%Y-%m-%d %H:%M:%S', '%Y-%m-%d %H:%M', '%Y-%m-%d', '%d/%m/%Y %H:%M:%S', '%d/%m/%Y'):
        try:
            return datetime.strptime(date_str, fmt)
        except ValueError:
            pass
    return None

def extract_line_timestamp(line_text):
    if not line_text:
        return None
    for pattern in TIMESTAMP_REGEX_PATTERNS:
        match = pattern.search(line_text)
        if match:
            ts_raw = match.group(1).replace('T', ' ')
            dt = parse_date_str(ts_raw)
            if dt:
                return dt
            try:
                curr_year = datetime.now().year
                return datetime.strptime(f"{curr_year} {ts_raw}", "%Y %b %d %H:%M:%S")
            except Exception:
                pass
    return None



@app.route('/log-analyzer')
@login_required
def log_analyzer():
    return render_template('log_analyzer.html')


@app.route('/api/logs/upload', methods=['POST'])
@login_required
def api_upload_log_file():
    if 'file' not in request.files:
        return jsonify({'success': False, 'message': 'No file submitted in request.'}), 400

    uploaded_files = request.files.getlist('file')
    if not uploaded_files or all(f.filename == '' for f in uploaded_files):
        return jsonify({'success': False, 'message': 'No file selected.'}), 400

    saved_files = []

    conn = get_db()
    cursor = conn.cursor()

    for file in uploaded_files:
        if not file or not file.filename:
            continue

        orig_filename = secure_filename(file.filename) or file.filename
        ext = os.path.splitext(orig_filename)[1].lower()
        if not ext:
            ext = '.log'

        stored_filename = f"{uuid.uuid4().hex}{ext}"
        target_path = os.path.join(UPLOAD_LOGS_DIR, stored_filename)
        file.save(target_path)

        file_size = os.path.getsize(target_path)
        lines = read_log_file_lines(target_path)
        line_count = len(lines)
        uploaded_by = session.get('username') or session.get('full_name') or 'User'

        cursor.execute('''
            INSERT INTO uploaded_log_files (filename, stored_filename, file_type, file_size, line_count, uploaded_by)
            VALUES (?, ?, ?, ?, ?, ?);
        ''', (orig_filename, stored_filename, ext.lstrip('.').upper(), file_size, line_count, uploaded_by))
        
        file_id = cursor.lastrowid
        saved_files.append({
            'id': file_id,
            'filename': orig_filename,
            'stored_filename': stored_filename,
            'file_type': ext.lstrip('.').upper(),
            'file_size': file_size,
            'line_count': line_count,
            'uploaded_by': uploaded_by,
            'created_at': datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        })

    conn.commit()
    conn.close()

    if saved_files:
        log_activity('UPLOAD_LOG_FILE', 'uploaded_log_files', f"{len(saved_files)} file(s)", f"Uploaded: {', '.join(f['filename'] for f in saved_files)}")

    return jsonify({'success': True, 'message': f'Successfully uploaded {len(saved_files)} log file(s).', 'files': saved_files})


@app.route('/api/logs/files')
@login_required
def api_get_log_files():
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM uploaded_log_files ORDER BY created_at DESC;")
    rows = [dict(r) for r in cursor.fetchall()]
    conn.close()
    return jsonify({'success': True, 'files': rows})


@app.route('/api/logs/delete/<int:file_id>', methods=['POST'])
@login_required
def api_delete_log_file(file_id):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM uploaded_log_files WHERE id = ?;", (file_id,))
    row = cursor.fetchone()

    if not row:
        conn.close()
        return jsonify({'success': False, 'message': 'Log file record not found.'}), 404

    target_path = os.path.join(UPLOAD_LOGS_DIR, row['stored_filename'])
    if os.path.exists(target_path):
        try:
            os.remove(target_path)
        except Exception as e:
            print(f"Warning: Could not remove file {target_path}: {e}")

    cursor.execute("DELETE FROM uploaded_log_files WHERE id = ?;", (file_id,))
    conn.commit()
    conn.close()

    log_activity('DELETE_LOG_FILE', 'uploaded_log_files', str(file_id), f"Deleted log file: {row['filename']}")

    return jsonify({'success': True, 'message': f"Log file '{row['filename']}' deleted successfully."})


@app.route('/api/logs/search')
@login_required
def api_search_log_files():
    query = request.args.get('q', '').strip()
    file_id = request.args.get('file_id', 'all').strip()
    start_date_str = request.args.get('start_date', '').strip()
    end_date_str = request.args.get('end_date', '').strip()
    max_results = request.args.get('max_results', 500, type=int)

    start_dt = parse_date_str(start_date_str)
    end_dt = parse_date_str(end_date_str)

    conn = get_db()
    cursor = conn.cursor()

    if file_id != 'all' and file_id:
        f_ids = [int(i.strip()) for i in file_id.split(',') if i.strip().isdigit()]
        if f_ids:
            placeholders = ','.join(['?'] * len(f_ids))
            cursor.execute(f"SELECT * FROM uploaded_log_files WHERE id IN ({placeholders});", f_ids)
        else:
            cursor.execute("SELECT * FROM uploaded_log_files ORDER BY created_at DESC;")
    else:
        cursor.execute("SELECT * FROM uploaded_log_files ORDER BY created_at DESC;")

    file_records = cursor.fetchall()
    conn.close()

    if not file_records:
        return jsonify({'success': True, 'query': query, 'total_matches': 0, 'matches': [], 'entities': {'ips': [], 'vlans': []}})

    matches = []
    
    is_vlan_only = query.isdigit() and len(query) <= 5
    if is_vlan_only:
        q_pattern = re.compile(rf'\b(?:vlan[\s:=_-]*)?{re.escape(query)}\b', re.IGNORECASE)
    else:
        q_pattern = re.compile(re.escape(query) if query else r'.*', re.IGNORECASE)

    for f_rec in file_records:
        if len(matches) >= max_results:
            break

        f_path = os.path.join(UPLOAD_LOGS_DIR, f_rec['stored_filename'])
        lines = read_log_file_lines(f_path)

        for line_no, line_text in enumerate(lines, start=1):
            if len(matches) >= max_results:
                break

            if start_dt or end_dt:
                line_dt = extract_line_timestamp(line_text)
                if line_dt:
                    if start_dt and line_dt < start_dt:
                        continue
                    if end_dt and line_dt > end_dt:
                        continue

            if not query or q_pattern.search(line_text):
                matches.append({
                    'file_id': f_rec['id'],
                    'filename': f_rec['filename'],
                    'line_no': line_no,
                    'line_text': line_text
                })

    entities = extract_entities_from_matches(matches)

    return jsonify({
        'success': True,
        'query': query,
        'file_id': file_id,
        'start_date': start_date_str,
        'end_date': end_date_str,
        'total_matches': len(matches),
        'matches': matches,
        'entities': entities
    })


@app.route('/api/logs/export-search')
@login_required
def api_export_log_search():
    query = request.args.get('q', '').strip()
    file_id = request.args.get('file_id', 'all').strip()
    start_date_str = request.args.get('start_date', '').strip()
    end_date_str = request.args.get('end_date', '').strip()
    export_format = request.args.get('format', 'txt').strip().lower()

    start_dt = parse_date_str(start_date_str)
    end_dt = parse_date_str(end_date_str)

    conn = get_db()
    cursor = conn.cursor()

    if file_id != 'all' and file_id:
        f_ids = [int(i.strip()) for i in file_id.split(',') if i.strip().isdigit()]
        if f_ids:
            placeholders = ','.join(['?'] * len(f_ids))
            cursor.execute(f"SELECT * FROM uploaded_log_files WHERE id IN ({placeholders});", f_ids)
        else:
            cursor.execute("SELECT * FROM uploaded_log_files ORDER BY created_at DESC;")
    else:
        cursor.execute("SELECT * FROM uploaded_log_files ORDER BY created_at DESC;")

    file_records = cursor.fetchall()
    conn.close()

    matches = []
    q_pattern = re.compile(re.escape(query) if query else r'.*', re.IGNORECASE)

    for f_rec in file_records:
        f_path = os.path.join(UPLOAD_LOGS_DIR, f_rec['stored_filename'])
        lines = read_log_file_lines(f_path)

        for line_no, line_text in enumerate(lines, start=1):
            if start_dt or end_dt:
                line_dt = extract_line_timestamp(line_text)
                if line_dt:
                    if start_dt and line_dt < start_dt:
                        continue
                    if end_dt and line_dt > end_dt:
                        continue

            if not query or q_pattern.search(line_text):
                matches.append({
                    'filename': f_rec['filename'],
                    'line_no': line_no,
                    'line_text': line_text
                })

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    if export_format == 'csv':
        df = pd.DataFrame(matches)
        output = BytesIO()
        csv_bytes = df.to_csv(index=False, encoding='utf-8').encode('utf-8')
        output.write(csv_bytes)
        output.seek(0)
        return send_file(
            output,
            as_attachment=True,
            download_name=f"log_search_results_{timestamp}.csv",
            mimetype='text/csv'
        )
    else:
        output = BytesIO()
        content = "\n".join(f"[{m['filename']}:L{m['line_no']}] {m['line_text']}" for m in matches)
        output.write(content.encode('utf-8'))
        output.seek(0)
        return send_file(
            output,
            as_attachment=True,
            download_name=f"log_search_results_{timestamp}.txt",
            mimetype='text/plain'
        )



@app.route('/api/ping-cef-batch', methods=['POST'])
@login_required
def api_ping_cef_batch():
    data = request.get_json(silent=True) or {}
    raw_ips = data.get('ips', [])
    if not isinstance(raw_ips, list) or not raw_ips:
        return jsonify({'success': False, 'message': 'No IP list provided.'}), 400

    def ping_single_ip(raw_ip):
        if not raw_ip:
            return (raw_ip, {'status': 'DOWN', 'is_reachable': False, 'message': 'No IP'})

        clean_ip = str(raw_ip).split('/')[0].strip()
        try:
            ip_obj = str(ipaddress.ip_address(clean_ip))
        except ValueError:
            return (raw_ip, {'status': 'DOWN', 'is_reachable': False, 'message': 'Invalid IP format'})

        if os.name == 'nt':
            cmd = ['ping', '-n', '2', '-w', '1000', ip_obj]
        else:
            cmd = ['ping', '-c', '2', '-W', '1', ip_obj]

        try:
            proc = subprocess.run(cmd, capture_output=True, text=True, timeout=3.5)
            output = proc.stdout or proc.stderr or ''
            is_reachable = check_ip_reachability(ip_obj, proc.returncode, output)

            return (raw_ip, {
                'ip': ip_obj,
                'status': 'UP' if is_reachable else 'DOWN',
                'is_reachable': is_reachable
            })
        except Exception:
            return (raw_ip, {'ip': ip_obj, 'status': 'DOWN', 'is_reachable': False})

    unique_ips = list(set(raw_ips))
    results = {}
    with concurrent.futures.ThreadPoolExecutor(max_workers=min(30, max(1, len(unique_ips)))) as executor:
        futures = [executor.submit(ping_single_ip, ip) for ip in unique_ips]
        for future in concurrent.futures.as_completed(futures):
            r_ip, res = future.result()
            results[r_ip] = res

    return jsonify({'success': True, 'results': results})


@app.route('/site/new', methods=['GET', 'POST'])
@login_required
def create_site():
    if request.method == 'POST':
        site_id = request.form.get('site_id', '').strip()
        enodeb_address = request.form.get('enodeb_address', '').strip()
        site_name = request.form.get('site_name', '').strip()
        ssa = request.form.get('ssa', '').strip()
        location = request.form.get('Location', request.form.get('location', '')).strip()
        cpan_maan = request.form.get('cpan_maan_vsat', request.form.get('cpan/maan/vsat', '')).strip()
        
        oam_vlan = request.form.get('oam_vlan', '').strip()
        mgmt_rac_vlan = request.form.get('mgmt_rac_vlan', '').strip()
        s1_c_vlan = request.form.get('s1_c_vlan', '').strip()
        s1_u_vlan = request.form.get('s1_u_vlan', '').strip()
        mgmt_ip = request.form.get('mgmt_ip', '').strip()
        mgmt_gateway = request.form.get('mgmt_gateway', '').strip()
        s1_u_ip = request.form.get('s1_u_ip', '').strip()
        mme_ip = request.form.get('mme_ip', '').strip()
        endpoint_type = request.form.get('endpoint_type', '').strip()
        endpoint_node_router = request.form.get('endpoint_node_router', '').strip()
        endpoint_ip = request.form.get('endpoint_ip', '').strip()
        l3_gateway_maan = request.form.get('l3_gateway_maan', '').strip()
        endpoint_ports = transform_tx_port(request.form.get('endpoint_ports', '').strip())
        cpan_a_end_node = request.form.get('cpan_a_end_node', '').strip()
        cpan_a_end_ip = request.form.get('cpan_a_end_ip', '').strip()
        cpan_a_end_ports = transform_tx_port(request.form.get('cpan_a_end_ports', '').strip())
        cpan_z_end_node = request.form.get('cpan_z_end_node', '').strip()
        cpan_z_end_ip = request.form.get('cpan_z_end_ip', '').strip()
        cpan_service = request.form.get('cpan_service', '').strip()
        service_vlans = request.form.get('service_vlans', '').strip()
        maan_l3_interface = request.form.get('maan_l3_interface', '').strip()
        maan_vpn = request.form.get('maan_vpn', '').strip()
        mask = request.form.get('mask', '').strip()
        route_distinguisher = request.form.get('route_distinguisher', '').strip()
        as_num = request.form.get('as_num', '').strip()
        ems = request.form.get('ems', '').strip()
        oam_cef_ip_pool = request.form.get('oam_cef_ip_pool', '').strip()
        oam_hw_gw = request.form.get('oam_hw_gw', '').strip()
        oam_hw_ip = request.form.get('oam_hw_ip', '').strip()
        reason = request.form.get('reason', '').strip()
        
        tx_ip = endpoint_ip or cpan_a_end_ip or request.form.get('tx-system-ip', '').strip()
        tx_loc = endpoint_node_router or cpan_a_end_node or request.form.get('tx-system-location', '').strip()
        tx_port = endpoint_ports or cpan_a_end_ports or transform_tx_port(request.form.get('tx-system-port', '').strip())
        vlan = s1_c_vlan or service_vlans or oam_vlan or request.form.get('vlan', '').strip()
        
        if not site_id:
            flash('Site ID is required.', 'danger')
            return render_template('site_form.html', mode='create', form_data=request.form)
            
        conn = get_db()
        cursor = conn.cursor()
        
        cursor.execute("SELECT id FROM bts_sites WHERE LOWER(site_id) = LOWER(?);", (site_id,))
        if cursor.fetchone() is not None:
            conn.close()
            flash(f'Site ID "{site_id}" already exists. Please use a unique Site ID.', 'danger')
            return render_template('site_form.html', mode='create', form_data=request.form)
            
        if site_name and (not ssa or not location):
            derived_ssa, derived_loc = derive_ssa_and_location(site_name, ssa, location)
            if not ssa:
                ssa = derived_ssa
            if not location:
                location = derived_loc
                
        cursor.execute('''
            INSERT INTO bts_sites (
                enodeb_address, site_id, site_name, ssa, location, cpan_maan_vsat,
                oam_vlan, mgmt_rac_vlan, s1_c_vlan, s1_u_vlan, mgmt_ip, mgmt_gateway,
                s1_u_ip, mme_ip, endpoint_type, endpoint_node_router, endpoint_ip,
                l3_gateway_maan, endpoint_ports, cpan_a_end_node, cpan_a_end_ip,
                cpan_a_end_ports, cpan_z_end_node, cpan_z_end_ip, cpan_service,
                service_vlans, maan_l3_interface, maan_vpn, mask, route_distinguisher,
                as_num, ems, oam_cef_ip_pool, oam_hw_gw, oam_hw_ip,
                tx_system_ip, tx_system_location, tx_system_port, vlan, reason
            ) VALUES (
                ?, ?, ?, ?, ?, ?,
                ?, ?, ?, ?, ?, ?,
                ?, ?, ?, ?, ?,
                ?, ?, ?, ?,
                ?, ?, ?, ?,
                ?, ?, ?, ?, ?,
                ?, ?, ?, ?, ?,
                ?, ?, ?, ?, ?
            );
        ''', (
            enodeb_address, site_id, site_name, ssa, location, cpan_maan,
            oam_vlan, mgmt_rac_vlan, s1_c_vlan, s1_u_vlan, mgmt_ip, mgmt_gateway,
            s1_u_ip, mme_ip, endpoint_type, endpoint_node_router, endpoint_ip,
            l3_gateway_maan, endpoint_ports, cpan_a_end_node, cpan_a_end_ip,
            cpan_a_end_ports, cpan_z_end_node, cpan_z_end_ip, cpan_service,
            service_vlans, maan_l3_interface, maan_vpn, mask, route_distinguisher,
            as_num, ems, oam_cef_ip_pool, oam_hw_gw, oam_hw_ip,
            tx_ip, tx_loc, tx_port, vlan, reason
        ))
        
        conn.commit()
        conn.close()
        
        flash(f'Site "{site_id}" created successfully!', 'success')
        return redirect(url_for('view_site', site_id=site_id))
        
    return render_template('site_form.html', mode='create', form_data={})


@app.route('/site/<path:site_id>')
@login_required
def view_site(site_id):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM bts_sites WHERE LOWER(site_id) = LOWER(?);", (site_id,))
    row = cursor.fetchone()
    conn.close()
    
    if not row:
        flash(f'Site ID "{site_id}" not found.', 'warning')
        return redirect(url_for('index'))
        
    site = row_to_dict(row)
    return render_template('view_site.html', site=site)


@app.route('/site/<path:site_id>/edit', methods=['GET', 'POST'])
@login_required
def edit_site(site_id):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM bts_sites WHERE LOWER(site_id) = LOWER(?);", (site_id,))
    row = cursor.fetchone()
    
    if not row:
        conn.close()
        flash(f'Site ID "{site_id}" not found.', 'warning')
        return redirect(url_for('index'))
        
    site = row_to_dict(row)
    
    if request.method == 'POST':
        enodeb_address = request.form.get('enodeb_address', '').strip()
        site_name = request.form.get('site_name', '').strip()
        ssa = request.form.get('ssa', '').strip()
        location = request.form.get('Location', request.form.get('location', '')).strip()
        cpan_maan = request.form.get('cpan_maan_vsat', request.form.get('cpan/maan/vsat', '')).strip()
        
        oam_vlan = request.form.get('oam_vlan', '').strip()
        mgmt_rac_vlan = request.form.get('mgmt_rac_vlan', '').strip()
        s1_c_vlan = request.form.get('s1_c_vlan', '').strip()
        s1_u_vlan = request.form.get('s1_u_vlan', '').strip()
        mgmt_ip = request.form.get('mgmt_ip', '').strip()
        mgmt_gateway = request.form.get('mgmt_gateway', '').strip()
        s1_u_ip = request.form.get('s1_u_ip', '').strip()
        mme_ip = request.form.get('mme_ip', '').strip()
        endpoint_type = request.form.get('endpoint_type', '').strip()
        endpoint_node_router = request.form.get('endpoint_node_router', '').strip()
        endpoint_ip = request.form.get('endpoint_ip', '').strip()
        l3_gateway_maan = request.form.get('l3_gateway_maan', '').strip()
        endpoint_ports = transform_tx_port(request.form.get('endpoint_ports', '').strip())
        cpan_a_end_node = request.form.get('cpan_a_end_node', '').strip()
        cpan_a_end_ip = request.form.get('cpan_a_end_ip', '').strip()
        cpan_a_end_ports = transform_tx_port(request.form.get('cpan_a_end_ports', '').strip())
        cpan_z_end_node = request.form.get('cpan_z_end_node', '').strip()
        cpan_z_end_ip = request.form.get('cpan_z_end_ip', '').strip()
        cpan_service = request.form.get('cpan_service', '').strip()
        service_vlans = request.form.get('service_vlans', '').strip()
        maan_l3_interface = request.form.get('maan_l3_interface', '').strip()
        maan_vpn = request.form.get('maan_vpn', '').strip()
        mask = request.form.get('mask', '').strip()
        route_distinguisher = request.form.get('route_distinguisher', '').strip()
        as_num = request.form.get('as_num', '').strip()
        ems = request.form.get('ems', '').strip()
        oam_cef_ip_pool = request.form.get('oam_cef_ip_pool', '').strip()
        oam_hw_gw = request.form.get('oam_hw_gw', '').strip()
        oam_hw_ip = request.form.get('oam_hw_ip', '').strip()
        reason = request.form.get('reason', '').strip()
        
        tx_ip = endpoint_ip or cpan_a_end_ip or request.form.get('tx-system-ip', '').strip()
        tx_loc = endpoint_node_router or cpan_a_end_node or request.form.get('tx-system-location', '').strip()
        tx_port = endpoint_ports or cpan_a_end_ports or transform_tx_port(request.form.get('tx-system-port', '').strip())
        vlan = s1_c_vlan or service_vlans or oam_vlan or request.form.get('vlan', '').strip()
        
        if site_name and (not ssa or not location):
            derived_ssa, derived_loc = derive_ssa_and_location(site_name, ssa, location)
            if not ssa:
                ssa = derived_ssa
            if not location:
                location = derived_loc
                
        cursor.execute('''
            UPDATE bts_sites SET
                enodeb_address = ?,
                site_name = ?,
                ssa = ?,
                location = ?,
                cpan_maan_vsat = ?,
                oam_vlan = ?,
                mgmt_rac_vlan = ?,
                s1_c_vlan = ?,
                s1_u_vlan = ?,
                mgmt_ip = ?,
                mgmt_gateway = ?,
                s1_u_ip = ?,
                mme_ip = ?,
                endpoint_type = ?,
                endpoint_node_router = ?,
                endpoint_ip = ?,
                l3_gateway_maan = ?,
                endpoint_ports = ?,
                cpan_a_end_node = ?,
                cpan_a_end_ip = ?,
                cpan_a_end_ports = ?,
                cpan_z_end_node = ?,
                cpan_z_end_ip = ?,
                cpan_service = ?,
                service_vlans = ?,
                maan_l3_interface = ?,
                maan_vpn = ?,
                mask = ?,
                route_distinguisher = ?,
                as_num = ?,
                ems = ?,
                oam_cef_ip_pool = ?,
                oam_hw_gw = ?,
                oam_hw_ip = ?,
                tx_system_ip = ?,
                tx_system_location = ?,
                tx_system_port = ?,
                vlan = ?,
                reason = ?,
                updated_at = CURRENT_TIMESTAMP
            WHERE LOWER(site_id) = LOWER(?);
        ''', (
            enodeb_address, site_name, ssa, location, cpan_maan,
            oam_vlan, mgmt_rac_vlan, s1_c_vlan, s1_u_vlan, mgmt_ip, mgmt_gateway,
            s1_u_ip, mme_ip, endpoint_type, endpoint_node_router, endpoint_ip,
            l3_gateway_maan, endpoint_ports, cpan_a_end_node, cpan_a_end_ip,
            cpan_a_end_ports, cpan_z_end_node, cpan_z_end_ip, cpan_service,
            service_vlans, maan_l3_interface, maan_vpn, mask, route_distinguisher,
            as_num, ems, oam_cef_ip_pool, oam_hw_gw, oam_hw_ip,
            tx_ip, tx_loc, tx_port, vlan, reason, site_id
        ))
        
        conn.commit()
        conn.close()
        
        log_activity('EDIT_SITE', 'bts_sites', site_id, f'Updated site "{site_id}" ({site_name or ""})')
        flash(f'Site "{site_id}" updated successfully!', 'success')
        return redirect(url_for('view_site', site_id=site_id))
        
    conn.close()
    return render_template('site_form.html', mode='edit', site=site, form_data=site)


@app.route('/site/<path:site_id>/delete', methods=['POST'])
@login_required
def delete_site(site_id):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM bts_sites WHERE LOWER(site_id) = LOWER(?);", (site_id,))
    conn.commit()
    deleted = cursor.rowcount > 0
    conn.close()
    
    if deleted:
        flash(f'Site "{site_id}" deleted successfully!', 'success')
    else:
        flash(f'Site ID "{site_id}" not found.', 'warning')
        
    return redirect(url_for('index'))


ALL_REPORT_COLUMNS = {
    'site_id': 'Site Id',
    'site_name': 'Site Name',
    'status': 'Status (UP/DOWN)',
    'enodeb_address': 'eNodeB Address',
    'ssa': 'SSA',
    'location': 'Location',
    'cpan_maan_vsat': 'MAAN/CPAN/VSAT',
    'oam_vlan': 'OAM VLAN',
    'mgmt_rac_vlan': 'Mgmt / RAC VLAN',
    's1_c_vlan': 'S1-C VLAN',
    's1_u_vlan': 'S1-U VLAN',
    'mgmt_ip': 'Mgmt IP',
    'mgmt_gateway': 'Mgmt Gateway',
    's1_u_ip': 'S1-U IP',
    'mme_ip': 'MME IP',
    'endpoint_type': 'Endpoint Type',
    'endpoint_node_router': 'Endpoint — Node / Router',
    'endpoint_ip': 'Endpoint IP',
    'l3_gateway_maan': 'L3 Gateway (MAAN)',
    'endpoint_ports': 'Endpoint Port(s)',
    'cpan_a_end_node': 'CPAN A End Node',
    'cpan_a_end_ip': 'CPAN A End IP',
    'cpan_a_end_ports': 'CPAN A End Port(s)',
    'cpan_z_end_node': 'CPAN Z End Node',
    'cpan_z_end_ip': 'CPAN Z End IP',
    'cpan_service': 'CPAN Service',
    'service_vlans': 'Service VLANs',
    'maan_l3_interface': 'MAAN L3 Interface',
    'maan_vpn': 'MAAN VPN',
    'mask': 'Mask',
    'route_distinguisher': 'Route Distinguisher',
    'as_num': 'AS',
    'ems': 'EMS',
    'oam_cef_ip_pool': 'OAM CEF IP pool',
    'oam_hw_gw': 'OAM HW GW',
    'oam_hw_ip': 'OAM HW IP',
    'tx_system_ip': 'tx-system-ip',
    'tx_system_location': 'tx-system-location',
    'tx_system_port': 'tx-system-port',
    'vlan': 'VLAN',
    'reason': 'Reason / Remark'
}


def compute_ping_statuses_for_dataframe(df):
    """
    Given a pandas DataFrame containing site or node records, pings target IPs concurrently
    and populates a 'Status (UP/DOWN)' column.
    """
    if df is None or df.empty:
        if df is not None:
            df['Status (UP/DOWN)'] = []
        return df

    ip_candidates = []
    for _, row in df.iterrows():
        raw_ip = (row.get('_temp_target_ip') or
                  row.get('Endpoint IP') or row.get('endpoint_ip') or
                  row.get('NE IP') or row.get('ne_ip') or
                  row.get('CPAN A End IP') or row.get('cpan_a_end_ip') or
                  row.get('tx-system-ip') or row.get('tx_system_ip') or
                  row.get('eNodeB Address') or row.get('enodeb_address') or
                  row.get('Mgmt IP') or row.get('mgmt_ip') or '')
        clean_ip = str(raw_ip).split('/')[0].strip() if raw_ip else ''
        ip_candidates.append(clean_ip)

    unique_ips = list(set([ip for ip in ip_candidates if ip and ip.lower() not in ('nan', 'none', '-', '')]))
    ping_map = {}

    def ping_single(ip_str):
        try:
            ip_obj = str(ipaddress.ip_address(ip_str))
        except ValueError:
            return (ip_str, 'N/A')

        cmd = ['ping', '-n', '1', '-w', '800', ip_obj] if os.name == 'nt' else ['ping', '-c', '1', '-W', '1', ip_obj]
        try:
            proc = subprocess.run(cmd, capture_output=True, text=True, timeout=1.5)
            output = proc.stdout or proc.stderr or ''
            is_up = check_ip_reachability(ip_obj, proc.returncode, output)
            return (ip_str, 'UP' if is_up else 'DOWN')
        except Exception:
            return (ip_str, 'DOWN')

    if unique_ips:
        with concurrent.futures.ThreadPoolExecutor(max_workers=min(50, max(1, len(unique_ips)))) as executor:
            futures = [executor.submit(ping_single, ip) for ip in unique_ips]
            for future in concurrent.futures.as_completed(futures):
                ip, status = future.result()
                ping_map[ip] = status

    statuses = []
    for ip in ip_candidates:
        if not ip or ip.lower() in ('nan', 'none', '-', ''):
            statuses.append('N/A')
        else:
            statuses.append(ping_map.get(ip, 'DOWN'))

    df['Status (UP/DOWN)'] = statuses
    return df


def format_openpyxl_report(file_stream_or_wb):
    """
    Applies styling for headers, gridlines, and UP/DOWN status column in Excel export.
    """
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

    if isinstance(file_stream_or_wb, (str, BytesIO)):
        file_stream_or_wb.seek(0)
        wb = openpyxl.load_workbook(file_stream_or_wb)
    else:
        wb = file_stream_or_wb

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")

    up_fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
    up_font = Font(name="Segoe UI", size=10, bold=True, color="155724")

    down_fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
    down_font = Font(name="Segoe UI", size=10, bold=True, color="721C24")

    na_font = Font(name="Segoe UI", size=10, color="6C757D")

    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    for sheet in wb.worksheets:
        sheet.views.sheetView[0].showGridLines = True

        status_col_idx = None
        for col_idx in range(1, sheet.max_column + 1):
            cell = sheet.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            val = str(cell.value or '')
            if 'Status' in val or 'UP/DOWN' in val:
                status_col_idx = col_idx

        for row in range(2, sheet.max_row + 1):
            for col in range(1, sheet.max_column + 1):
                cell = sheet.cell(row=row, column=col)
                cell.border = thin_border
                cell.font = Font(name="Segoe UI", size=9.5)

                if col == status_col_idx:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    s_val = str(cell.value or '').strip().upper()
                    if s_val == 'UP':
                        cell.fill = up_fill
                        cell.font = up_font
                    elif s_val == 'DOWN':
                        cell.fill = down_fill
                        cell.font = down_font
                    else:
                        cell.font = na_font

        # Auto-adjust column widths
        for col in sheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out


@app.route('/export')
@login_required
def export_excel():
    conn = get_db()
    export_sql = '''
        SELECT
            enodeb_address AS 'eNodeB Address',
            site_id AS 'Site Id',
            site_name AS 'Site Name',
            cpan_maan_vsat AS 'MAAN/CPAN/VSAT',
            oam_vlan AS 'OAM VLAN',
            mgmt_rac_vlan AS 'Mgmt / RAC VLAN',
            s1_c_vlan AS 'S1-C VLAN',
            s1_u_vlan AS 'S1-U VLAN',
            mgmt_ip AS 'Mgmt IP',
            mgmt_gateway AS 'Mgmt Gateway',
            s1_u_ip AS 'S1-U IP',
            mme_ip AS 'MME IP',
            endpoint_type AS 'Endpoint Type',
            endpoint_node_router AS 'Endpoint — Node / Router',
            endpoint_ip AS 'Endpoint IP',
            l3_gateway_maan AS 'L3 Gateway (MAAN)',
            endpoint_ports AS 'Endpoint Port(s)',
            cpan_a_end_node AS 'CPAN A End Node',
            cpan_a_end_ip AS 'CPAN A End IP',
            cpan_a_end_ports AS 'CPAN A End Port(s)',
            cpan_z_end_node AS 'CPAN Z End Node',
            cpan_z_end_ip AS 'CPAN Z End IP',
            cpan_service AS 'CPAN Service',
            service_vlans AS 'Service VLANs',
            maan_l3_interface AS 'MAAN L3 Interface',
            maan_vpn AS 'MAAN VPN',
            mask AS 'Mask',
            route_distinguisher AS 'Route Distinguisher',
            as_num AS 'AS',
            ems AS 'EMS',
            oam_cef_ip_pool AS 'OAM CEF IP pool',
            oam_hw_gw AS 'OAM HW GW',
            oam_hw_ip AS 'OAM HW IP'
        FROM bts_sites ORDER BY id ASC;
    '''
    df = pd.read_sql_query(export_sql, conn)
    conn.close()
    
    if not df.empty:
        df = compute_ping_statuses_for_dataframe(df)
        if 'Status (UP/DOWN)' in df.columns and 'Site Name' in df.columns:
            cols = list(df.columns)
            cols.remove('Status (UP/DOWN)')
            idx = cols.index('Site Name') + 1
            cols.insert(idx, 'Status (UP/DOWN)')
            df = df[cols]
            
    output = BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Sheet1')
    output = format_openpyxl_report(output)
    
    return send_file(
        output,
        as_attachment=True,
        download_name='btsdatabase_updated.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )


@app.route('/export/custom', methods=['GET', 'POST'])
@login_required
def export_custom_report():
    if request.method == 'POST':
        export_format = request.form.get('format', 'xlsx').lower()
        scope = request.form.get('scope', 'filtered')
        query = request.form.get('q', '').strip()
        search_by = request.form.get('search_by', 'all').strip()
        selected_ssa = request.form.get('ssa', '').strip()
        sort_by = request.form.get('sort_by', 'id').strip()
        sort_order = request.form.get('sort_order', 'asc').strip()
        selected_cols = request.form.getlist('cols')
    else:
        export_format = request.args.get('format', 'xlsx').lower()
        scope = request.args.get('scope', 'filtered')
        query = request.args.get('q', '').strip()
        search_by = request.args.get('search_by', 'all').strip()
        selected_ssa = request.args.get('ssa', '').strip()
        sort_by = request.args.get('sort_by', 'id').strip()
        sort_order = request.args.get('sort_order', 'asc').strip()
        selected_cols = request.args.getlist('cols')
        
    if not selected_cols:
        selected_cols = list(ALL_REPORT_COLUMNS.keys())
    else:
        selected_cols = [c for c in selected_cols if c in ALL_REPORT_COLUMNS]
        if not selected_cols:
            selected_cols = list(ALL_REPORT_COLUMNS.keys())
            
    db_cols = [c for c in selected_cols if c != 'status']
    select_parts = [f"{col} AS '{ALL_REPORT_COLUMNS[col]}'" for col in db_cols]
    
    if 'status' in selected_cols:
        for ip_c in ['endpoint_ip', 'cpan_a_end_ip', 'tx_system_ip', 'enodeb_address']:
            if ip_c not in db_cols:
                select_parts.append(f"{ip_c} AS '_temp_{ip_c}'")
                
    select_sql = ", ".join(select_parts)
    
    sort_col = VALID_SORT_COLUMNS.get(sort_by, 'bts_sites.id')
    sort_dir = "DESC" if str(sort_order).lower() in ("desc", "descending") else "ASC"
    
    conn = get_db()
    cursor = conn.cursor()
    items = parse_search_query(query)
    
    if scope == 'filtered' and len(items) > 1:
        if search_by not in VALID_SEARCH_COLUMNS:
            search_by = "all"
            
        ssa_where = ""
        ssa_params = []
        if selected_ssa:
            ssa_where = "WHERE LOWER(ssa) = LOWER(?)"
            ssa_params = [selected_ssa.strip()]
            
        if search_by != "all" and search_by in ALL_SEARCHABLE_COLS:
            text_expr = f"IFNULL({search_by}, '')"
        else:
            text_expr = " || ' ' || ".join([f"IFNULL({c}, '')" for c in ALL_SEARCHABLE_COLS])
            
        sql = f"SELECT id, ({text_expr}) as search_text FROM bts_sites {ssa_where} ORDER BY {sort_col} {sort_dir};"
        cursor.execute(sql, ssa_params)
        all_rows = cursor.fetchall()
        
        pattern = re.compile('|'.join([re.escape(it) for it in items]), re.IGNORECASE)
        matched_ids = [r['id'] for r in all_rows if pattern.search(r['search_text'])]
        
        if matched_ids:
            placeholders = ','.join(['?']*len(matched_ids))
            query_sql = f"SELECT {select_sql} FROM bts_sites WHERE id IN ({placeholders}) ORDER BY {sort_col} {sort_dir};"
            df = pd.read_sql_query(query_sql, conn, params=matched_ids)
        else:
            df = pd.DataFrame(columns=[ALL_REPORT_COLUMNS[c] for c in selected_cols if c in ALL_REPORT_COLUMNS])
    else:
        where_clauses = []
        params = []
        
        if scope == 'filtered':
            if search_by not in VALID_SEARCH_COLUMNS:
                search_by = "all"
                
            if query:
                q_like = f"%{query}%"
                if search_by != "all":
                    where_clauses.append(f"LOWER({search_by}) LIKE LOWER(?)")
                    params.append(q_like)
                else:
                    or_clauses = [f"LOWER({c}) LIKE LOWER(?)" for c in ALL_SEARCHABLE_COLS]
                    where_clauses.append("(" + " OR ".join(or_clauses) + ")")
                    params.extend([q_like] * len(ALL_SEARCHABLE_COLS))
                    
            if selected_ssa:
                where_clauses.append("LOWER(ssa) = LOWER(?)")
                params.append(selected_ssa)
                
        where_sql = ""
        if where_clauses:
            where_sql = "WHERE " + " AND ".join(where_clauses)
            
        query_sql = f"SELECT {select_sql} FROM bts_sites {where_sql} ORDER BY {sort_col} {sort_dir};"
        df = pd.read_sql_query(query_sql, conn, params=params)
    conn.close()
    
    if 'status' in selected_cols and not df.empty:
        ip_series = []
        for _, row in df.iterrows():
            ep = (row.get('Endpoint IP') or row.get('_temp_endpoint_ip') or
                  row.get('CPAN A End IP') or row.get('_temp_cpan_a_end_ip') or
                  row.get('tx-system-ip') or row.get('_temp_tx_system_ip') or
                  row.get('eNodeB Address') or row.get('_temp_enodeb_address') or '')
            ip_series.append(ep)
        df['_temp_target_ip'] = ip_series
        df = compute_ping_statuses_for_dataframe(df)
        
    # Drop temp columns
    for c in list(df.columns):
        if str(c).startswith('_temp_'):
            df.drop(columns=[c], inplace=True)
            
    # Order columns as requested in selected_cols
    desired_headers = [ALL_REPORT_COLUMNS[c] for c in selected_cols if ALL_REPORT_COLUMNS[c] in df.columns]
    if desired_headers:
        df = df[desired_headers]
        
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    
    if export_format == 'csv':
        filename = f"CPAN_Report_{timestamp}.csv"
        output = BytesIO()
        csv_bytes = df.to_csv(index=False, encoding='utf-8').encode('utf-8')
        output.write(csv_bytes)
        output.seek(0)
        return send_file(
            output,
            as_attachment=True,
            download_name=filename,
            mimetype='text/csv'
        )
    else:
        filename = f"CPAN_Report_{timestamp}.xlsx"
        output = BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='Custom Report')
        output = format_openpyxl_report(output)
        return send_file(
            output,
            as_attachment=True,
            download_name=filename,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )


# ==========================================
# CPAN NODES MANAGEMENT & SEARCH ROUTES
# ==========================================

VALID_CPAN_SEARCH_COLUMNS = {
    'all': 'Generic (All Columns)',
    'ne_ip': 'NE IP',
    'location': 'Location',
    'type': 'Type',
    'ssa': 'SSA',
    'phase': 'Phase',
    'ne_name': 'NE Name',
    'dcc_ip': 'DCC IP'
}

ALL_CPAN_SEARCHABLE_COLS = [
    'ne_ip', 'location', 'type', 'ssa', 'phase', 'ne_name',
    'dcc_ip', 'orig_location'
]


VALID_CPAN_SORT_COLUMNS = {
    'id': 'cpan_nodes.id',
    'ne_ip': 'ne_ip',
    'location': 'location',
    'type': 'type',
    'ssa': 'ssa',
    'phase': 'phase',
    'ne_name': 'ne_name',
    'dcc_ip': 'dcc_ip'
}


def search_cpan_nodes(query="", search_by="all", selected_ssa="", selected_type="", page=1, per_page=25, sort_by="id", sort_order="asc"):
    conn = get_db()
    cursor = conn.cursor()
    
    if search_by not in VALID_CPAN_SEARCH_COLUMNS:
        search_by = "all"
        
    sort_col = VALID_CPAN_SORT_COLUMNS.get(sort_by, 'cpan_nodes.id')
    sort_dir = "DESC" if str(sort_order).lower() in ("desc", "descending") else "ASC"
    
    cursor.execute("SELECT COUNT(*) FROM cpan_nodes;")
    total_records = cursor.fetchone()[0]
    
    cursor.execute("SELECT DISTINCT ssa FROM cpan_nodes WHERE ssa IS NOT NULL AND ssa != '' ORDER BY ssa ASC;")
    unique_ssas = [r['ssa'] for r in cursor.fetchall()]
    
    cursor.execute("SELECT DISTINCT type FROM cpan_nodes WHERE type IS NOT NULL AND type != '' ORDER BY type ASC;")
    unique_types = [r['type'] for r in cursor.fetchall()]
    
    cursor.execute("SELECT COUNT(DISTINCT location) FROM cpan_nodes WHERE location IS NOT NULL AND location != '';")
    unique_locations_count = cursor.fetchone()[0]
    
    items = parse_search_query(query)
    
    if len(items) > 1:
        where_parts = []
        params = []
        if selected_ssa:
            where_parts.append("LOWER(ssa) = LOWER(?)")
            params.append(selected_ssa.strip())
        if selected_type:
            where_parts.append("LOWER(type) = LOWER(?)")
            params.append(selected_type.strip())
            
        where_sql = ("WHERE " + " AND ".join(where_parts)) if where_parts else ""
        
        if search_by != "all" and search_by in ALL_CPAN_SEARCHABLE_COLS:
            text_expr = f"IFNULL({search_by}, '')"
        else:
            text_expr = " || ' ' || ".join([f"IFNULL({c}, '')" for c in ALL_CPAN_SEARCHABLE_COLS])
            
        sql = f"SELECT id, ({text_expr}) as search_text FROM cpan_nodes {where_sql} ORDER BY {sort_col} {sort_dir};"
        cursor.execute(sql, params)
        all_rows = cursor.fetchall()
        
        pattern = re.compile('|'.join([re.escape(it) for it in items]), re.IGNORECASE)
        matched_ids = [r['id'] for r in all_rows if pattern.search(r['search_text'])]
        
        filtered_count = len(matched_ids)
        total_pages = max(1, (filtered_count + per_page - 1) // per_page)
        page = max(1, min(page, total_pages))
        offset = (page - 1) * per_page
        
        page_ids = matched_ids[offset : offset + per_page]
        
        if page_ids:
            placeholders = ','.join(['?']*len(page_ids))
            cursor.execute(f"SELECT * FROM cpan_nodes WHERE id IN ({placeholders}) ORDER BY {sort_col} {sort_dir};", page_ids)
            rows = [dict(r) for r in cursor.fetchall()]
        else:
            rows = []
    else:
        where_clauses = []
        params = []
        
        if query:
            q_like = f"%{query.strip()}%"
            if search_by != "all" and search_by in ALL_CPAN_SEARCHABLE_COLS:
                where_clauses.append(f"LOWER({search_by}) LIKE LOWER(?)")
                params.append(q_like)
            else:
                or_clauses = [f"LOWER({c}) LIKE LOWER(?)" for c in ALL_CPAN_SEARCHABLE_COLS]
                where_clauses.append("(" + " OR ".join(or_clauses) + ")")
                params.extend([q_like] * len(ALL_CPAN_SEARCHABLE_COLS))
                
        if selected_ssa:
            where_clauses.append("LOWER(ssa) = LOWER(?)")
            params.append(selected_ssa.strip())
            
        if selected_type:
            where_clauses.append("LOWER(type) = LOWER(?)")
            params.append(selected_type.strip())
            
        where_sql = ""
        if where_clauses:
            where_sql = "WHERE " + " AND ".join(where_clauses)
            
        count_sql = f"SELECT COUNT(*) FROM cpan_nodes {where_sql};"
        cursor.execute(count_sql, params)
        filtered_count = cursor.fetchone()[0]
        
        total_pages = max(1, (filtered_count + per_page - 1) // per_page)
        page = max(1, min(page, total_pages))
        offset = (page - 1) * per_page
        
        data_sql = f"SELECT * FROM cpan_nodes {where_sql} ORDER BY {sort_col} {sort_dir} LIMIT ? OFFSET ?;"
        cursor.execute(data_sql, params + [per_page, offset])
        rows = [dict(r) for r in cursor.fetchall()]
        
    conn.close()
    
    return {
        'nodes': rows,
        'rows': rows,
        'total_records': total_records,
        'filtered_count': filtered_count,
        'unique_ssas': unique_ssas,
        'unique_types': unique_types,
        'unique_locations_count': unique_locations_count,
        'page': page,
        'per_page': per_page,
        'total_pages': total_pages,
        'query': query,
        'search_by': search_by,
        'selected_ssa': selected_ssa,
        'selected_type': selected_type,
        'sort_by': sort_by,
        'sort_order': sort_order
    }


@app.route('/cpan-nodes')
@login_required
def cpan_nodes():
    query = request.args.get('query', '').strip()
    search_by = request.args.get('search_by', 'all').strip()
    selected_ssa = request.args.get('ssa', '').strip()
    selected_type = request.args.get('type', '').strip()
    sort_by = request.args.get('sort_by', 'id').strip()
    sort_order = request.args.get('sort_order', 'asc').strip()
    
    try:
        page = int(request.args.get('page', 1))
    except ValueError:
        page = 1
        
    try:
        per_page = int(request.args.get('per_page', 25))
        if per_page not in (10, 25, 50, 100, 250, 500, 1000, 2000):
            per_page = 25
    except ValueError:
        per_page = 25
        
    res = search_cpan_nodes(query, search_by, selected_ssa, selected_type, page, per_page, sort_by=sort_by, sort_order=sort_order)
    
    return render_template(
        'cpan_nodes.html',
        nodes=res['rows'],
        total_records=res['total_records'],
        filtered_count=res['filtered_count'],
        unique_ssas=res['unique_ssas'],
        unique_types=res['unique_types'],
        unique_locations_count=res['unique_locations_count'],
        page=res['page'],
        per_page=res['per_page'],
        total_pages=res['total_pages'],
        query=query,
        search_by=search_by,
        selected_ssa=selected_ssa,
        selected_type=selected_type,
        sort_by=sort_by,
        sort_order=sort_order,
        search_columns=VALID_CPAN_SEARCH_COLUMNS
    )


@app.route('/api/cpan-nodes/search')
@login_required
def api_search_cpan_nodes():
    query = request.args.get('query', '').strip()
    search_by = request.args.get('search_by', 'all').strip()
    selected_ssa = request.args.get('ssa', '').strip()
    selected_type = request.args.get('type', '').strip()
    sort_by = request.args.get('sort_by', 'id').strip()
    sort_order = request.args.get('sort_order', 'asc').strip()
    
    try:
        page = int(request.args.get('page', 1))
    except ValueError:
        page = 1
        
    try:
        per_page = int(request.args.get('per_page', 25))
        if per_page not in (10, 25, 50, 100, 250, 500, 1000, 2000):
            per_page = 25
    except ValueError:
        per_page = 25
        
    res = search_cpan_nodes(query, search_by, selected_ssa, selected_type, page, per_page, sort_by=sort_by, sort_order=sort_order)
    return jsonify(res)



@app.route('/cpan-node/new', methods=['POST'])
@login_required
def create_cpan_node():
    ne_ip = request.form.get('ne_ip', '').strip()
    location = request.form.get('location', '').strip()
    ne_type = request.form.get('type', '').strip()
    ssa = request.form.get('ssa', '').strip()
    phase = request.form.get('phase', '').strip()
    ne_name = request.form.get('ne_name', '').strip()
    dcc_ip = request.form.get('dcc_ip', '').strip()
    
    if not ne_ip or not location:
        flash('NE IP and Location are required fields.', 'warning')
        return redirect(url_for('cpan_nodes'))
        
    if not ne_name:
        ne_name = f"{ne_ip}_{location}_{ne_type}_{ssa}_{phase}".rstrip('_')
        
    conn = get_db()
    cursor = conn.cursor()
    try:
        cursor.execute('''
            INSERT INTO cpan_nodes (ne_ip, location, type, ssa, phase, ne_name, dcc_ip)
            VALUES (?, ?, ?, ?, ?, ?, ?);
        ''', (ne_ip, location, ne_type, ssa, phase, ne_name, dcc_ip))
        conn.commit()
        flash(f'CPAN Node {ne_ip} ({location}) added successfully!', 'success')
    except Exception as e:
        conn.rollback()
        flash(f'Error adding CPAN Node: {str(e)}', 'danger')
    finally:
        conn.close()
        
    return redirect(url_for('cpan_nodes'))


@app.route('/api/cpan-node/<int:node_id>')
@login_required
def get_cpan_node(node_id):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM cpan_nodes WHERE id = ?;", (node_id,))
    row = cursor.fetchone()
    conn.close()
    
    if row is None:
        return jsonify({'error': 'CPAN Node not found'}), 404
        
    return jsonify(dict(row))


@app.route('/cpan-node/<int:node_id>/edit', methods=['POST'])
@login_required
def edit_cpan_node(node_id):
    ne_ip = request.form.get('ne_ip', '').strip()
    location = request.form.get('location', '').strip()
    ne_type = request.form.get('type', '').strip()
    ssa = request.form.get('ssa', '').strip()
    phase = request.form.get('phase', '').strip()
    ne_name = request.form.get('ne_name', '').strip()
    dcc_ip = request.form.get('dcc_ip', '').strip()
    
    if not ne_ip or not location:
        flash('NE IP and Location are required fields.', 'warning')
        return redirect(url_for('cpan_nodes'))
        
    conn = get_db()
    cursor = conn.cursor()
    try:
        cursor.execute('''
            UPDATE cpan_nodes
            SET ne_ip = ?, location = ?, type = ?, ssa = ?, phase = ?, ne_name = ?, dcc_ip = ?
            WHERE id = ?;
        ''', (ne_ip, location, ne_type, ssa, phase, ne_name, dcc_ip, node_id))
        conn.commit()
        flash(f'CPAN Node {ne_ip} ({location}) updated successfully!', 'success')
    except Exception as e:
        conn.rollback()
        flash(f'Error updating CPAN Node: {str(e)}', 'danger')
    finally:
        conn.close()
        
    return redirect(url_for('cpan_nodes'))


@app.route('/cpan-node/<int:node_id>/delete', methods=['POST'])
@login_required
def delete_cpan_node(node_id):
    conn = get_db()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT ne_ip, location FROM cpan_nodes WHERE id = ?;", (node_id,))
        node = cursor.fetchone()
        node_label = f"{node['ne_ip']} ({node['location']})" if node else f"ID #{node_id}"
            
        cursor.execute("DELETE FROM cpan_nodes WHERE id = ?;", (node_id,))
        conn.commit()
        flash(f'CPAN Node {node_label} deleted successfully!', 'success')
    except Exception as e:
        conn.rollback()
        flash(f'Error deleting CPAN Node: {str(e)}', 'danger')
    finally:
        conn.close()
        
    return redirect(url_for('cpan_nodes'))


@app.route('/upload-cpan-nodes', methods=['POST'])
@login_required
def upload_cpan_nodes():
    if 'file' not in request.files:
        flash('No file uploaded.', 'danger')
        return redirect(url_for('cpan_nodes'))
        
    file = request.files['file']
    if file.filename == '':
        flash('No file selected.', 'warning')
        return redirect(url_for('cpan_nodes'))
        
    if not (file.filename.endswith('.csv') or file.filename.endswith('.xlsx') or file.filename.endswith('.xls')):
        flash('Invalid file format. Please upload a CSV file (e.g. GUJ_CPAN-NODE_LIST.csv).', 'danger')
        return redirect(url_for('cpan_nodes'))
        
    try:
        if file.filename.endswith('.csv'):
            df = pd.read_csv(file)
        else:
            df = pd.read_excel(file)
            
        from init_db import import_cpan_nodes_csv
        imported_count = import_cpan_nodes_csv(df)
        flash(f'Successfully uploaded and processed {imported_count} CPAN Nodes into database!', 'success')
    except Exception as e:
        flash(f'Error processing file: {str(e)}', 'danger')
        
    return redirect(url_for('cpan_nodes'))


@app.route('/export-cpan-nodes')
@login_required
def export_cpan_nodes():
    query = request.args.get('query', '').strip()
    search_by = request.args.get('search_by', 'all').strip()
    selected_ssa = request.args.get('ssa', '').strip()
    selected_type = request.args.get('type', '').strip()
    sort_by = request.args.get('sort_by', 'id').strip()
    sort_order = request.args.get('sort_order', 'asc').strip()
    export_format = request.args.get('format', 'csv').strip().lower()
    
    sort_col = VALID_CPAN_SORT_COLUMNS.get(sort_by, 'cpan_nodes.id')
    sort_dir = "DESC" if str(sort_order).lower() in ("desc", "descending") else "ASC"
    
    conn = get_db()
    cursor = conn.cursor()
    
    items = parse_search_query(query)
    
    if len(items) > 1:
        where_parts = []
        params = []
        if selected_ssa:
            where_parts.append("LOWER(ssa) = LOWER(?)")
            params.append(selected_ssa.strip())
        if selected_type:
            where_parts.append("LOWER(type) = LOWER(?)")
            params.append(selected_type.strip())
            
        where_sql = ("WHERE " + " AND ".join(where_parts)) if where_parts else ""
        
        if search_by != "all" and search_by in ALL_CPAN_SEARCHABLE_COLS:
            text_expr = f"IFNULL({search_by}, '')"
        else:
            text_expr = " || ' ' || ".join([f"IFNULL({c}, '')" for c in ALL_CPAN_SEARCHABLE_COLS])
            
        sql = f"SELECT id, ({text_expr}) as search_text FROM cpan_nodes {where_sql} ORDER BY {sort_col} {sort_dir};"
        cursor.execute(sql, params)
        all_rows = cursor.fetchall()
        
        pattern = re.compile('|'.join([re.escape(it) for it in items]), re.IGNORECASE)
        matched_ids = [r['id'] for r in all_rows if pattern.search(r['search_text'])]
        
        if matched_ids:
            placeholders = ','.join(['?']*len(matched_ids))
            query_sql = f"""
                SELECT ne_ip AS 'NE IP', location AS 'Location', type AS 'Type', ssa AS 'SSA', phase AS 'Phase',
                       ne_name AS 'NE Name', dcc_ip AS 'DCC IP'
                FROM cpan_nodes WHERE id IN ({placeholders}) ORDER BY {sort_col} {sort_dir};
            """
            df = pd.read_sql_query(query_sql, conn, params=matched_ids)
        else:
            df = pd.DataFrame(columns=['NE IP', 'Location', 'Type', 'SSA', 'Phase', 'NE Name', 'DCC IP'])
    else:
        where_clauses = []
        params = []
        
        if search_by not in VALID_CPAN_SEARCH_COLUMNS:
            search_by = "all"
            
        if query:
            q_like = f"%{query.strip()}%"
            if search_by != "all" and search_by in ALL_CPAN_SEARCHABLE_COLS:
                where_clauses.append(f"LOWER({search_by}) LIKE LOWER(?)")
                params.append(q_like)
            else:
                or_clauses = [f"LOWER({c}) LIKE LOWER(?)" for c in ALL_CPAN_SEARCHABLE_COLS]
                where_clauses.append("(" + " OR ".join(or_clauses) + ")")
                params.extend([q_like] * len(ALL_CPAN_SEARCHABLE_COLS))
                
        if selected_ssa:
            where_clauses.append("LOWER(ssa) = LOWER(?)")
            params.append(selected_ssa)
            
        if selected_type:
            where_clauses.append("LOWER(type) = LOWER(?)")
            params.append(selected_type)
            
        where_sql = ""
        if where_clauses:
            where_sql = "WHERE " + " AND ".join(where_clauses)
            
        query_sql = f"""
            SELECT ne_ip AS 'NE IP', location AS 'Location', type AS 'Type', ssa AS 'SSA', phase AS 'Phase',
                   ne_name AS 'NE Name', dcc_ip AS 'DCC IP'
            FROM cpan_nodes {where_sql} ORDER BY {sort_col} {sort_dir};
        """
        df = pd.read_sql_query(query_sql, conn, params=params)
        
    conn.close()
    
    if not df.empty:
        df = compute_ping_statuses_for_dataframe(df)
        if 'Status (UP/DOWN)' in df.columns and 'NE IP' in df.columns:
            cols = list(df.columns)
            cols.remove('Status (UP/DOWN)')
            idx = cols.index('NE IP') + 1
            cols.insert(idx, 'Status (UP/DOWN)')
            df = df[cols]
            
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    
    if export_format in ('excel', 'xlsx'):
        filename = f"CPAN_Nodes_{timestamp}.xlsx"
        output = BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='CPAN Nodes')
        output = format_openpyxl_report(output)
        return send_file(
            output,
            as_attachment=True,
            download_name=filename,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
    else:
        filename = f"CPAN_Nodes_{timestamp}.csv"
        output = BytesIO()
        csv_bytes = df.to_csv(index=False, encoding='utf-8').encode('utf-8')
        output.write(csv_bytes)
        output.seek(0)
        return send_file(
            output,
            as_attachment=True,
            download_name=filename,
            mimetype='text/csv'
        )


# ==========================================
# CPAN DL LIST & SERVICES ROUTES
# ==========================================

VALID_CPAN_DL_SEARCH_COLUMNS = {
    'all': 'Generic (All Columns)',
    'name': 'Name',
    'media_type': 'Media Type',
    'bandwidth': 'Bandwidth',
    'signal_type': 'Signal Type',
    'a_end': 'A End',
    'z_end': 'Z End',
    'alarm_status': 'Alarm Status',
    'client': 'Client'
}

ALL_CPAN_DL_SEARCHABLE_COLS = [
    'name', 'media_type', 'bandwidth', 'signal_type', 'direction',
    'a_end', 'z_end', 'alarm_status', 'order_name', 'creator', 'client', 'description'
]

VALID_CPAN_DL_SORT_COLUMNS = {
    'id': 'cpan_dl_list.id',
    'name': 'name',
    'media_type': 'media_type',
    'bandwidth': 'bandwidth',
    'signal_type': 'signal_type',
    'a_end': 'a_end',
    'z_end': 'z_end',
    'alarm_status': 'alarm_status'
}


def search_cpan_dl_list(query="", search_by="all", selected_media="", page=1, per_page=25, sort_by="id", sort_order="asc"):
    conn = get_db()
    cursor = conn.cursor()
    
    if search_by not in VALID_CPAN_DL_SEARCH_COLUMNS:
        search_by = "all"
        
    sort_col = VALID_CPAN_DL_SORT_COLUMNS.get(sort_by, 'cpan_dl_list.id')
    sort_dir = "DESC" if str(sort_order).lower() in ("desc", "descending") else "ASC"
    
    cursor.execute("SELECT COUNT(*) FROM cpan_dl_list;")
    total_records = cursor.fetchone()[0]
    
    cursor.execute("SELECT DISTINCT media_type FROM cpan_dl_list WHERE media_type IS NOT NULL AND media_type != '' ORDER BY media_type ASC;")
    unique_media_types = [r['media_type'] for r in cursor.fetchall()]
    
    where_clauses = []
    params = []
    
    if query:
        q_like = f"%{query.strip()}%"
        if search_by != "all" and search_by in ALL_CPAN_DL_SEARCHABLE_COLS:
            where_clauses.append(f"LOWER({search_by}) LIKE LOWER(?)")
            params.append(q_like)
        else:
            or_clauses = [f"LOWER({c}) LIKE LOWER(?)" for c in ALL_CPAN_DL_SEARCHABLE_COLS]
            where_clauses.append("(" + " OR ".join(or_clauses) + ")")
            params.extend([q_like] * len(ALL_CPAN_DL_SEARCHABLE_COLS))
            
    if selected_media:
        where_clauses.append("LOWER(media_type) = LOWER(?)")
        params.append(selected_media.strip())
        
    where_sql = ""
    if where_clauses:
        where_sql = "WHERE " + " AND ".join(where_clauses)
        
    count_sql = f"SELECT COUNT(*) FROM cpan_dl_list {where_sql};"
    cursor.execute(count_sql, params)
    filtered_count = cursor.fetchone()[0]
    
    total_pages = max(1, (filtered_count + per_page - 1) // per_page)
    page = max(1, min(page, total_pages))
    offset = (page - 1) * per_page
    
    data_sql = f"SELECT * FROM cpan_dl_list {where_sql} ORDER BY {sort_col} {sort_dir} LIMIT ? OFFSET ?;"
    cursor.execute(data_sql, params + [per_page, offset])
    rows = [dict(r) for r in cursor.fetchall()]
    
    conn.close()
    
    return {
        'rows': rows,
        'total_records': total_records,
        'filtered_count': filtered_count,
        'unique_media_types': unique_media_types,
        'page': page,
        'per_page': per_page,
        'total_pages': total_pages,
        'query': query,
        'search_by': search_by,
        'selected_media': selected_media,
        'sort_by': sort_by,
        'sort_order': sort_order
    }


@app.route('/cpan-dl-list')
@login_required
def cpan_dl_list():
    query = request.args.get('query', '').strip()
    search_by = request.args.get('search_by', 'all').strip()
    selected_media = request.args.get('media', '').strip()
    sort_by = request.args.get('sort_by', 'id').strip()
    sort_order = request.args.get('sort_order', 'asc').strip()
    
    try:
        page = int(request.args.get('page', 1))
    except ValueError:
        page = 1
        
    try:
        per_page = int(request.args.get('per_page', 25))
        if per_page not in (10, 25, 50, 100, 250, 500, 1000, 2000):
            per_page = 25
    except ValueError:
        per_page = 25
        
    res = search_cpan_dl_list(query, search_by, selected_media, page, per_page, sort_by=sort_by, sort_order=sort_order)
    
    return render_template(
        'cpan_dl_list.html',
        rows=res['rows'],
        total_records=res['total_records'],
        filtered_count=res['filtered_count'],
        unique_media_types=res['unique_media_types'],
        page=res['page'],
        per_page=res['per_page'],
        total_pages=res['total_pages'],
        query=query,
        search_by=search_by,
        selected_media=selected_media,
        sort_by=sort_by,
        sort_order=sort_order,
        search_columns=VALID_CPAN_DL_SEARCH_COLUMNS
    )


@app.route('/export-cpan-dl-list')
@login_required
def export_cpan_dl_list():
    query = request.args.get('query', '').strip()
    search_by = request.args.get('search_by', 'all').strip()
    selected_media = request.args.get('media', '').strip()
    sort_by = request.args.get('sort_by', 'id').strip()
    sort_order = request.args.get('sort_order', 'asc').strip()
    export_format = request.args.get('format', 'csv').strip().lower()
    
    sort_col = VALID_CPAN_DL_SORT_COLUMNS.get(sort_by, 'cpan_dl_list.id')
    sort_dir = "DESC" if str(sort_order).lower() in ("desc", "descending") else "ASC"
    
    conn = get_db()
    
    where_clauses = []
    params = []
    
    if search_by not in VALID_CPAN_DL_SEARCH_COLUMNS:
        search_by = "all"
        
    if query:
        q_like = f"%{query.strip()}%"
        if search_by != "all" and search_by in ALL_CPAN_DL_SEARCHABLE_COLS:
            where_clauses.append(f"LOWER({search_by}) LIKE LOWER(?)")
            params.append(q_like)
        else:
            or_clauses = [f"LOWER({c}) LIKE LOWER(?)" for c in ALL_CPAN_DL_SEARCHABLE_COLS]
            where_clauses.append("(" + " OR ".join(or_clauses) + ")")
            params.extend([q_like] * len(ALL_CPAN_DL_SEARCHABLE_COLS))
            
    if selected_media:
        where_clauses.append("LOWER(media_type) = LOWER(?)")
        params.append(selected_media)
        
    where_sql = ""
    if where_clauses:
        where_sql = "WHERE " + " AND ".join(where_clauses)
        
    query_sql = f"""
        SELECT name AS 'Name', media_type AS 'Media Type', bandwidth AS 'Bandwidth',
               signal_type AS 'Signal Type', direction AS 'Direction', a_end AS 'A End',
               z_end AS 'Z End', alarm_status AS 'Alarm Status', cir_utilization AS 'CIR Util %',
               bandwidth_utilization AS 'BW Util %', order_name AS 'Order Name',
               creator AS 'Creator', client AS 'Client', create_time AS 'Create Time',
               update_time AS 'Update Time', description AS 'Description'
        FROM cpan_dl_list {where_sql} ORDER BY {sort_col} {sort_dir};
    """
    df = pd.read_sql_query(query_sql, conn, params=params)
    conn.close()
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    
    if export_format in ('excel', 'xlsx'):
        filename = f"CPAN_DL_List_{timestamp}.xlsx"
        output = BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='CPAN DL List')
        output.seek(0)
        return send_file(
            output,
            as_attachment=True,
            download_name=filename,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
    else:
        filename = f"CPAN_DL_List_{timestamp}.csv"
        output = BytesIO()
        csv_bytes = df.to_csv(index=False, encoding='utf-8').encode('utf-8')
        output.write(csv_bytes)
        output.seek(0)
        return send_file(
            output,
            as_attachment=True,
            download_name=filename,
            mimetype='text/csv'
        )


VALID_CPAN_SERVICES_SEARCH_COLUMNS = {
    'all': 'Generic (All Columns)',
    'name': 'Name',
    'service_type': 'Service Type',
    'order_name': 'Order Name',
    'client': 'Client',
    'a_end': 'A End',
    'z_end': 'Z End'
}

ALL_CPAN_SERVICES_SEARCHABLE_COLS = [
    'name', 'service_type', 'cos', 'trust_ce_qos', 'order_name',
    'client', 'a_end', 'z_end', 'description'
]

VALID_CPAN_SERVICES_SORT_COLUMNS = {
    'id': 'cpan_services.id',
    'name': 'name',
    'service_type': 'service_type',
    'bandwidth_kbps': 'bandwidth_kbps',
    'traffic_cir_kbps': 'traffic_cir_kbps',
    'order_name': 'order_name',
    'client': 'client',
    'a_end': 'a_end',
    'z_end': 'z_end'
}


def search_cpan_services(query="", search_by="all", selected_type="", page=1, per_page=25, sort_by="id", sort_order="asc"):
    conn = get_db()
    cursor = conn.cursor()
    
    if search_by not in VALID_CPAN_SERVICES_SEARCH_COLUMNS:
        search_by = "all"
        
    sort_col = VALID_CPAN_SERVICES_SORT_COLUMNS.get(sort_by, 'cpan_services.id')
    sort_dir = "DESC" if str(sort_order).lower() in ("desc", "descending") else "ASC"
    
    cursor.execute("SELECT COUNT(*) FROM cpan_services;")
    total_records = cursor.fetchone()[0]
    
    cursor.execute("SELECT DISTINCT service_type FROM cpan_services WHERE service_type IS NOT NULL AND service_type != '' ORDER BY service_type ASC;")
    unique_service_types = [r['service_type'] for r in cursor.fetchall()]
    
    where_clauses = []
    params = []
    
    if query:
        q_like = f"%{query.strip()}%"
        if search_by != "all" and search_by in ALL_CPAN_SERVICES_SEARCHABLE_COLS:
            where_clauses.append(f"LOWER({search_by}) LIKE LOWER(?)")
            params.append(q_like)
        else:
            or_clauses = [f"LOWER({c}) LIKE LOWER(?)" for c in ALL_CPAN_SERVICES_SEARCHABLE_COLS]
            where_clauses.append("(" + " OR ".join(or_clauses) + ")")
            params.extend([q_like] * len(ALL_CPAN_SERVICES_SEARCHABLE_COLS))
            
    if selected_type:
        where_clauses.append("LOWER(service_type) = LOWER(?)")
        params.append(selected_type.strip())
        
    where_sql = ""
    if where_clauses:
        where_sql = "WHERE " + " AND ".join(where_clauses)
        
    count_sql = f"SELECT COUNT(*) FROM cpan_services {where_sql};"
    cursor.execute(count_sql, params)
    filtered_count = cursor.fetchone()[0]
    
    total_pages = max(1, (filtered_count + per_page - 1) // per_page)
    page = max(1, min(page, total_pages))
    offset = (page - 1) * per_page
    
    data_sql = f"SELECT * FROM cpan_services {where_sql} ORDER BY {sort_col} {sort_dir} LIMIT ? OFFSET ?;"
    cursor.execute(data_sql, params + [per_page, offset])
    rows = [dict(r) for r in cursor.fetchall()]
    
    conn.close()
    
    return {
        'rows': rows,
        'total_records': total_records,
        'filtered_count': filtered_count,
        'unique_service_types': unique_service_types,
        'page': page,
        'per_page': per_page,
        'total_pages': total_pages,
        'query': query,
        'search_by': search_by,
        'selected_type': selected_type,
        'sort_by': sort_by,
        'sort_order': sort_order
    }


@app.route('/cpan-services')
@login_required
def cpan_services():
    query = request.args.get('query', '').strip()
    search_by = request.args.get('search_by', 'all').strip()
    selected_type = request.args.get('type', '').strip()
    sort_by = request.args.get('sort_by', 'id').strip()
    sort_order = request.args.get('sort_order', 'asc').strip()
    
    try:
        page = int(request.args.get('page', 1))
    except ValueError:
        page = 1
        
    try:
        per_page = int(request.args.get('per_page', 25))
        if per_page not in (10, 25, 50, 100, 250, 500, 1000, 2000):
            per_page = 25
    except ValueError:
        per_page = 25
        
    res = search_cpan_services(query, search_by, selected_type, page, per_page, sort_by=sort_by, sort_order=sort_order)
    
    return render_template(
        'cpan_services.html',
        rows=res['rows'],
        total_records=res['total_records'],
        filtered_count=res['filtered_count'],
        unique_service_types=res['unique_service_types'],
        page=res['page'],
        per_page=res['per_page'],
        total_pages=res['total_pages'],
        query=query,
        search_by=search_by,
        selected_type=selected_type,
        sort_by=sort_by,
        sort_order=sort_order,
        search_columns=VALID_CPAN_SERVICES_SEARCH_COLUMNS
    )


@app.route('/export-cpan-services')
@login_required
def export_cpan_services():
    query = request.args.get('query', '').strip()
    search_by = request.args.get('search_by', 'all').strip()
    selected_type = request.args.get('type', '').strip()
    sort_by = request.args.get('sort_by', 'id').strip()
    sort_order = request.args.get('sort_order', 'asc').strip()
    export_format = request.args.get('format', 'csv').strip().lower()
    
    sort_col = VALID_CPAN_SERVICES_SORT_COLUMNS.get(sort_by, 'cpan_services.id')
    sort_dir = "DESC" if str(sort_order).lower() in ("desc", "descending") else "ASC"
    
    conn = get_db()
    
    where_clauses = []
    params = []
    
    if search_by not in VALID_CPAN_SERVICES_SEARCH_COLUMNS:
        search_by = "all"
        
    if query:
        q_like = f"%{query.strip()}%"
        if search_by != "all" and search_by in ALL_CPAN_SERVICES_SEARCHABLE_COLS:
            where_clauses.append(f"LOWER({search_by}) LIKE LOWER(?)")
            params.append(q_like)
        else:
            or_clauses = [f"LOWER({c}) LIKE LOWER(?)" for c in ALL_CPAN_SERVICES_SEARCHABLE_COLS]
            where_clauses.append("(" + " OR ".join(or_clauses) + ")")
            params.extend([q_like] * len(ALL_CPAN_SERVICES_SEARCHABLE_COLS))
            
    if selected_type:
        where_clauses.append("LOWER(service_type) = LOWER(?)")
        params.append(selected_type)
        
    where_sql = ""
    if where_clauses:
        where_sql = "WHERE " + " AND ".join(where_clauses)
        
    query_sql = f"""
        SELECT name AS 'Name', service_type AS 'Service Type', bandwidth_kbps AS 'Bandwidth (Kbps)',
               traffic_cir_kbps AS 'Traffic CIR (Kbps)', traffic_eir_kbps AS 'Traffic EIR (Kbps)',
               network_cir_kbps AS 'Network CIR (Kbps)', network_eir_kbps AS 'Network EIR (Kbps)',
               cos AS 'Cos', trust_ce_qos AS 'Trust CE QoS', order_name AS 'Order Name',
               client AS 'Client', a_end AS 'A End', z_end AS 'Z End', create_time AS 'Create Time',
               update_time AS 'Update Time', description AS 'Description'
        FROM cpan_services {where_sql} ORDER BY {sort_col} {sort_dir};
    """
    df = pd.read_sql_query(query_sql, conn, params=params)
    conn.close()
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    
    if export_format in ('excel', 'xlsx'):
        filename = f"CPAN_Services_{timestamp}.xlsx"
        output = BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='CPAN Services')
        output.seek(0)
        return send_file(
            output,
            as_attachment=True,
            download_name=filename,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
    else:
        filename = f"CPAN_Services_{timestamp}.csv"
        output = BytesIO()
        csv_bytes = df.to_csv(index=False, encoding='utf-8').encode('utf-8')
        output.write(csv_bytes)
        output.seek(0)
        return send_file(
            output,
            as_attachment=True,
            download_name=filename,
            mimetype='text/csv'
        )


# ==========================================
# CPAN DL LIST API & CRUD ROUTES
# ==========================================

@app.route('/api/cpan-dl-list/search')
@login_required
def api_search_cpan_dl_list():
    query = request.args.get('query', '').strip()
    search_by = request.args.get('search_by', 'all').strip()
    selected_media = request.args.get('media', '').strip()
    sort_by = request.args.get('sort_by', 'id').strip()
    sort_order = request.args.get('sort_order', 'asc').strip()
    
    try:
        page = int(request.args.get('page', 1))
    except ValueError:
        page = 1
        
    try:
        per_page = int(request.args.get('per_page', 25))
        if per_page not in (10, 25, 50, 100, 250, 500, 1000, 2000):
            per_page = 25
    except ValueError:
        per_page = 25
        
    res = search_cpan_dl_list(query, search_by, selected_media, page, per_page, sort_by=sort_by, sort_order=sort_order)
    return jsonify(res)


@app.route('/cpan-dl/new', methods=['POST'])
@login_required
def create_cpan_dl():
    name = request.form.get('name', '').strip()
    media_type = request.form.get('media_type', '').strip()
    bandwidth = request.form.get('bandwidth', '').strip()
    signal_type = request.form.get('signal_type', '').strip()
    direction = request.form.get('direction', 'BI-DIR').strip()
    a_end = request.form.get('a_end', '').strip()
    z_end = request.form.get('z_end', '').strip()
    alarm_status = request.form.get('alarm_status', '-').strip()
    cir_utilization = request.form.get('cir_utilization', '').strip()
    bandwidth_utilization = request.form.get('bandwidth_utilization', '').strip()
    order_name = request.form.get('order_name', '').strip()
    creator = request.form.get('creator', session.get('username', '')).strip()
    client = request.form.get('client', '').strip()
    cost = request.form.get('cost', '1.0').strip()
    description = request.form.get('description', '').strip()
    
    if not name:
        flash('Name is a required field for CPAN DL record.', 'warning')
        return redirect(url_for('cpan_dl_list'))
        
    conn = get_db()
    cursor = conn.cursor()
    try:
        cursor.execute('''
            INSERT INTO cpan_dl_list (
                name, media_type, bandwidth, signal_type, direction,
                a_end, z_end, alarm_status, cir_utilization, bandwidth_utilization,
                order_name, creator, client, cost, description
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
        ''', (
            name, media_type, bandwidth, signal_type, direction,
            a_end, z_end, alarm_status, cir_utilization, bandwidth_utilization,
            order_name, creator, client, cost, description
        ))
        conn.commit()
        flash(f'CPAN DL Record "{name}" added successfully!', 'success')
    except Exception as e:
        conn.rollback()
        flash(f'Error adding CPAN DL Record: {str(e)}', 'danger')
    finally:
        conn.close()
        
    return redirect(url_for('cpan_dl_list'))


@app.route('/api/cpan-dl/<int:dl_id>')
@login_required
def get_cpan_dl(dl_id):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM cpan_dl_list WHERE id = ?;", (dl_id,))
    row = cursor.fetchone()
    conn.close()
    
    if row is None:
        return jsonify({'error': 'CPAN DL Record not found'}), 404
        
    return jsonify(dict(row))


@app.route('/cpan-dl/<int:dl_id>/edit', methods=['POST'])
@login_required
def edit_cpan_dl(dl_id):
    name = request.form.get('name', '').strip()
    media_type = request.form.get('media_type', '').strip()
    bandwidth = request.form.get('bandwidth', '').strip()
    signal_type = request.form.get('signal_type', '').strip()
    direction = request.form.get('direction', 'BI-DIR').strip()
    a_end = request.form.get('a_end', '').strip()
    z_end = request.form.get('z_end', '').strip()
    alarm_status = request.form.get('alarm_status', '-').strip()
    cir_utilization = request.form.get('cir_utilization', '').strip()
    bandwidth_utilization = request.form.get('bandwidth_utilization', '').strip()
    order_name = request.form.get('order_name', '').strip()
    creator = request.form.get('creator', '').strip()
    client = request.form.get('client', '').strip()
    cost = request.form.get('cost', '1.0').strip()
    description = request.form.get('description', '').strip()
    
    if not name:
        flash('Name is required.', 'warning')
        return redirect(url_for('cpan_dl_list'))
        
    conn = get_db()
    cursor = conn.cursor()
    try:
        cursor.execute('''
            UPDATE cpan_dl_list
            SET name = ?, media_type = ?, bandwidth = ?, signal_type = ?, direction = ?,
                a_end = ?, z_end = ?, alarm_status = ?, cir_utilization = ?, bandwidth_utilization = ?,
                order_name = ?, creator = ?, client = ?, cost = ?, description = ?
            WHERE id = ?;
        ''', (
            name, media_type, bandwidth, signal_type, direction,
            a_end, z_end, alarm_status, cir_utilization, bandwidth_utilization,
            order_name, creator, client, cost, description, dl_id
        ))
        conn.commit()
        flash(f'CPAN DL Record #{dl_id} updated successfully!', 'success')
    except Exception as e:
        conn.rollback()
        flash(f'Error updating CPAN DL Record: {str(e)}', 'danger')
    finally:
        conn.close()
        
    return redirect(url_for('cpan_dl_list'))


@app.route('/cpan-dl/<int:dl_id>/delete', methods=['POST'])
@login_required
def delete_cpan_dl(dl_id):
    conn = get_db()
    cursor = conn.cursor()
    try:
        cursor.execute("DELETE FROM cpan_dl_list WHERE id = ?;", (dl_id,))
        conn.commit()
        flash(f'CPAN DL Record #{dl_id} deleted successfully!', 'success')
    except Exception as e:
        conn.rollback()
        flash(f'Error deleting CPAN DL Record: {str(e)}', 'danger')
    finally:
        conn.close()
        
    return redirect(url_for('cpan_dl_list'))


@app.route('/upload-cpan-dl-list', methods=['POST'])
@login_required
def upload_cpan_dl_list():
    if 'file' not in request.files:
        flash('No file uploaded.', 'danger')
        return redirect(url_for('cpan_dl_list'))
        
    file = request.files['file']
    if file.filename == '':
        flash('No file selected.', 'warning')
        return redirect(url_for('cpan_dl_list'))
        
    if not (file.filename.endswith('.csv') or file.filename.endswith('.xlsx') or file.filename.endswith('.xls')):
        flash('Invalid file format. Please upload a CSV or Excel file (e.g. CPAN_DL_LIST.csv).', 'danger')
        return redirect(url_for('cpan_dl_list'))
        
    try:
        if file.filename.endswith('.csv'):
            df = pd.read_csv(file)
        else:
            df = pd.read_excel(file)
            
        from init_db import import_cpan_dl_list_csv
        imported_count = import_cpan_dl_list_csv(df)
        flash(f'Successfully uploaded and processed {imported_count} CPAN DL records into database!', 'success')
    except Exception as e:
        flash(f'Error processing file: {str(e)}', 'danger')
        
    return redirect(url_for('cpan_dl_list'))


# ==========================================
# CPAN SERVICES API & CRUD ROUTES
# ==========================================

@app.route('/api/cpan-services/search')
@login_required
def api_search_cpan_services():
    query = request.args.get('query', '').strip()
    search_by = request.args.get('search_by', 'all').strip()
    selected_type = request.args.get('type', '').strip()
    sort_by = request.args.get('sort_by', 'id').strip()
    sort_order = request.args.get('sort_order', 'asc').strip()
    
    try:
        page = int(request.args.get('page', 1))
    except ValueError:
        page = 1
        
    try:
        per_page = int(request.args.get('per_page', 25))
        if per_page not in (10, 25, 50, 100, 250, 500, 1000, 2000):
            per_page = 25
    except ValueError:
        per_page = 25
        
    res = search_cpan_services(query, search_by, selected_type, page, per_page, sort_by=sort_by, sort_order=sort_order)
    return jsonify(res)


@app.route('/cpan-service/new', methods=['POST'])
@login_required
def create_cpan_service():
    name = request.form.get('name', '').strip()
    service_type = request.form.get('service_type', '').strip()
    bandwidth_kbps = request.form.get('bandwidth_kbps', '0.0').strip()
    traffic_cir_kbps = request.form.get('traffic_cir_kbps', '0.0').strip()
    traffic_eir_kbps = request.form.get('traffic_eir_kbps', '0.0').strip()
    network_cir_kbps = request.form.get('network_cir_kbps', '0.0').strip()
    network_eir_kbps = request.form.get('network_eir_kbps', '0.0').strip()
    cos = request.form.get('cos', '').strip()
    trust_ce_qos = request.form.get('trust_ce_qos', '').strip()
    order_name = request.form.get('order_name', '').strip()
    client = request.form.get('client', '').strip()
    a_end = request.form.get('a_end', '').strip()
    z_end = request.form.get('z_end', '').strip()
    description = request.form.get('description', '').strip()
    
    if not name:
        flash('Service Name is required.', 'warning')
        return redirect(url_for('cpan_services'))
        
    conn = get_db()
    cursor = conn.cursor()
    try:
        cursor.execute('''
            INSERT INTO cpan_services (
                name, service_type, bandwidth_kbps, traffic_cir_kbps, traffic_eir_kbps,
                network_cir_kbps, network_eir_kbps, cos, trust_ce_qos, order_name,
                client, a_end, z_end, description
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
        ''', (
            name, service_type, bandwidth_kbps, traffic_cir_kbps, traffic_eir_kbps,
            network_cir_kbps, network_eir_kbps, cos, trust_ce_qos, order_name,
            client, a_end, z_end, description
        ))
        conn.commit()
        flash(f'CPAN Service "{name}" added successfully!', 'success')
    except Exception as e:
        conn.rollback()
        flash(f'Error adding CPAN Service: {str(e)}', 'danger')
    finally:
        conn.close()
        
    return redirect(url_for('cpan_services'))


@app.route('/api/cpan-service/<int:service_id>')
@login_required
def get_cpan_service(service_id):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM cpan_services WHERE id = ?;", (service_id,))
    row = cursor.fetchone()
    conn.close()
    
    if row is None:
        return jsonify({'error': 'CPAN Service not found'}), 404
        
    return jsonify(dict(row))


@app.route('/cpan-service/<int:service_id>/edit', methods=['POST'])
@login_required
def edit_cpan_service(service_id):
    name = request.form.get('name', '').strip()
    service_type = request.form.get('service_type', '').strip()
    bandwidth_kbps = request.form.get('bandwidth_kbps', '0.0').strip()
    traffic_cir_kbps = request.form.get('traffic_cir_kbps', '0.0').strip()
    traffic_eir_kbps = request.form.get('traffic_eir_kbps', '0.0').strip()
    network_cir_kbps = request.form.get('network_cir_kbps', '0.0').strip()
    network_eir_kbps = request.form.get('network_eir_kbps', '0.0').strip()
    cos = request.form.get('cos', '').strip()
    trust_ce_qos = request.form.get('trust_ce_qos', '').strip()
    order_name = request.form.get('order_name', '').strip()
    client = request.form.get('client', '').strip()
    a_end = request.form.get('a_end', '').strip()
    z_end = request.form.get('z_end', '').strip()
    description = request.form.get('description', '').strip()
    
    if not name:
        flash('Service Name is required.', 'warning')
        return redirect(url_for('cpan_services'))
        
    conn = get_db()
    cursor = conn.cursor()
    try:
        cursor.execute('''
            UPDATE cpan_services
            SET name = ?, service_type = ?, bandwidth_kbps = ?, traffic_cir_kbps = ?,
                traffic_eir_kbps = ?, network_cir_kbps = ?, network_eir_kbps = ?, cos = ?,
                trust_ce_qos = ?, order_name = ?, client = ?, a_end = ?, z_end = ?, description = ?
            WHERE id = ?;
        ''', (
            name, service_type, bandwidth_kbps, traffic_cir_kbps, traffic_eir_kbps,
            network_cir_kbps, network_eir_kbps, cos, trust_ce_qos, order_name,
            client, a_end, z_end, description, service_id
        ))
        conn.commit()
        flash(f'CPAN Service #{service_id} updated successfully!', 'success')
    except Exception as e:
        conn.rollback()
        flash(f'Error updating CPAN Service: {str(e)}', 'danger')
    finally:
        conn.close()
        
    return redirect(url_for('cpan_services'))


@app.route('/cpan-service/<int:service_id>/delete', methods=['POST'])
@login_required
def delete_cpan_service(service_id):
    conn = get_db()
    cursor = conn.cursor()
    try:
        cursor.execute("DELETE FROM cpan_services WHERE id = ?;", (service_id,))
        conn.commit()
        flash(f'CPAN Service #{service_id} deleted successfully!', 'success')
    except Exception as e:
        conn.rollback()
        flash(f'Error deleting CPAN Service: {str(e)}', 'danger')
    finally:
        conn.close()
        
    return redirect(url_for('cpan_services'))


@app.route('/upload-cpan-services', methods=['POST'])
@login_required
def upload_cpan_services():
    if 'file' not in request.files:
        flash('No file uploaded.', 'danger')
        return redirect(url_for('cpan_services'))
        
    file = request.files['file']
    if file.filename == '':
        flash('No file selected.', 'warning')
        return redirect(url_for('cpan_services'))
        
    if not (file.filename.endswith('.csv') or file.filename.endswith('.xlsx') or file.filename.endswith('.xls')):
        flash('Invalid file format. Please upload a CSV or Excel file (e.g. CPAN_Service_List.csv).', 'danger')
        return redirect(url_for('cpan_services'))
        
    try:
        if file.filename.endswith('.csv'):
            df = pd.read_csv(file)
        else:
            df = pd.read_excel(file)
            
        from init_db import import_cpan_services_csv
        imported_count = import_cpan_services_csv(df)
        flash(f'Successfully uploaded and processed {imported_count} CPAN services into database!', 'success')
    except Exception as e:
        flash(f'Error processing file: {str(e)}', 'danger')
        
    return redirect(url_for('cpan_services'))


# ==========================================
# MAAN NETWORK SEARCH, CRUD & EXPORT ROUTES
# ==========================================

VALID_MAAN_SEARCH_COLUMNS = {
    'all': 'Generic (All Columns)',
    'ne_ip': 'NE IP',
    'location': 'Location',
    'type': 'Type',
    'ssa': 'SSA',
    'phase': 'Phase',
    'ne_name': 'NE Name',
    'dcc_ip': 'DCC IP'
}
ALL_MAAN_SEARCHABLE_COLS = ['ne_ip', 'location', 'type', 'ssa', 'phase', 'ne_name', 'dcc_ip']
VALID_MAAN_SORT_COLUMNS = {
    'id': 'maan_nodes.id',
    'ne_ip': 'ne_ip',
    'location': 'location',
    'type': 'type',
    'ssa': 'ssa',
    'phase': 'phase',
    'ne_name': 'ne_name',
    'dcc_ip': 'dcc_ip'
}

VALID_MAAN_SERVICES_SEARCH_COLUMNS = {
    'all': 'Generic (All Columns)',
    'name': 'Name',
    'service_type': 'Service Type',
    'order_name': 'Order Name',
    'client': 'Client',
    'a_end': 'A End',
    'z_end': 'Z End'
}
ALL_MAAN_SERVICES_SEARCHABLE_COLS = ['name', 'service_type', 'order_name', 'client', 'a_end', 'z_end', 'description']
VALID_MAAN_SERVICES_SORT_COLUMNS = {
    'id': 'maan_services.id',
    'name': 'name',
    'service_type': 'service_type',
    'bandwidth_kbps': 'bandwidth_kbps',
    'traffic_cir_kbps': 'traffic_cir_kbps',
    'order_name': 'order_name',
    'client': 'client',
    'a_end': 'a_end',
    'z_end': 'z_end'
}

VALID_MAAN_TLS_SEARCH_COLUMNS = {
    'all': 'Generic (All Columns)',
    'name': 'Name',
    'media_type': 'Media Type',
    'bandwidth': 'Bandwidth',
    'signal_type': 'Signal Type',
    'a_end': 'A End',
    'z_end': 'Z End',
    'alarm_status': 'Alarm Status',
    'client': 'Client'
}
ALL_MAAN_TLS_SEARCHABLE_COLS = ['name', 'media_type', 'bandwidth', 'signal_type', 'a_end', 'z_end', 'alarm_status', 'order_name', 'client', 'description']
VALID_MAAN_TLS_SORT_COLUMNS = {
    'id': 'maan_tls.id',
    'name': 'name',
    'media_type': 'media_type',
    'bandwidth': 'bandwidth',
    'signal_type': 'signal_type',
    'a_end': 'a_end',
    'z_end': 'z_end',
    'alarm_status': 'alarm_status'
}


def search_maan_nodes(query="", search_by="all", selected_ssa="", selected_type="", page=1, per_page=25, sort_by="id", sort_order="asc"):
    conn = get_db()
    cursor = conn.cursor()
    
    if search_by not in VALID_MAAN_SEARCH_COLUMNS:
        search_by = "all"
        
    sort_col = VALID_MAAN_SORT_COLUMNS.get(sort_by, 'maan_nodes.id')
    sort_dir = "DESC" if str(sort_order).lower() in ("desc", "descending") else "ASC"
    
    cursor.execute("SELECT COUNT(*) FROM maan_nodes;")
    total_records = cursor.fetchone()[0]
    
    cursor.execute("SELECT DISTINCT ssa FROM maan_nodes WHERE ssa IS NOT NULL AND ssa != '' ORDER BY ssa ASC;")
    unique_ssas = [r['ssa'] for r in cursor.fetchall()]
    
    cursor.execute("SELECT DISTINCT type FROM maan_nodes WHERE type IS NOT NULL AND type != '' ORDER BY type ASC;")
    unique_types = [r['type'] for r in cursor.fetchall()]
    
    cursor.execute("SELECT COUNT(DISTINCT location) FROM maan_nodes WHERE location IS NOT NULL AND location != '';")
    unique_locations_count = cursor.fetchone()[0]
    
    where_clauses = []
    params = []
    
    if query:
        q_like = f"%{query.strip()}%"
        if search_by != "all" and search_by in ALL_MAAN_SEARCHABLE_COLS:
            where_clauses.append(f"LOWER({search_by}) LIKE LOWER(?)")
            params.append(q_like)
        else:
            or_clauses = [f"LOWER({c}) LIKE LOWER(?)" for c in ALL_MAAN_SEARCHABLE_COLS]
            where_clauses.append("(" + " OR ".join(or_clauses) + ")")
            params.extend([q_like] * len(ALL_MAAN_SEARCHABLE_COLS))
            
    if selected_ssa:
        where_clauses.append("LOWER(ssa) = LOWER(?)")
        params.append(selected_ssa.strip())
        
    if selected_type:
        where_clauses.append("LOWER(type) = LOWER(?)")
        params.append(selected_type.strip())
        
    where_sql = ""
    if where_clauses:
        where_sql = "WHERE " + " AND ".join(where_clauses)
        
    count_sql = f"SELECT COUNT(*) FROM maan_nodes {where_sql};"
    cursor.execute(count_sql, params)
    filtered_count = cursor.fetchone()[0]
    
    total_pages = max(1, (filtered_count + per_page - 1) // per_page)
    page = max(1, min(page, total_pages))
    offset = (page - 1) * per_page
    
    data_sql = f"SELECT * FROM maan_nodes {where_sql} ORDER BY {sort_col} {sort_dir} LIMIT ? OFFSET ?;"
    cursor.execute(data_sql, params + [per_page, offset])
    rows = [dict(r) for r in cursor.fetchall()]
    
    conn.close()
    
    return {
        'nodes': rows,
        'rows': rows,
        'total_records': total_records,
        'filtered_count': filtered_count,
        'unique_ssas': unique_ssas,
        'unique_types': unique_types,
        'unique_locations_count': unique_locations_count,
        'page': page,
        'per_page': per_page,
        'total_pages': total_pages,
        'query': query,
        'search_by': search_by,
        'selected_ssa': selected_ssa,
        'selected_type': selected_type,
        'sort_by': sort_by,
        'sort_order': sort_order
    }


def search_maan_services(query="", search_by="all", selected_type="", page=1, per_page=25, sort_by="id", sort_order="asc"):
    conn = get_db()
    cursor = conn.cursor()
    
    if search_by not in VALID_MAAN_SERVICES_SEARCH_COLUMNS:
        search_by = "all"
        
    sort_col = VALID_MAAN_SERVICES_SORT_COLUMNS.get(sort_by, 'maan_services.id')
    sort_dir = "DESC" if str(sort_order).lower() in ("desc", "descending") else "ASC"
    
    cursor.execute("SELECT COUNT(*) FROM maan_services;")
    total_records = cursor.fetchone()[0]
    
    cursor.execute("SELECT DISTINCT service_type FROM maan_services WHERE service_type IS NOT NULL AND service_type != '' ORDER BY service_type ASC;")
    unique_service_types = [r['service_type'] for r in cursor.fetchall()]
    
    where_clauses = []
    params = []
    
    if query:
        q_like = f"%{query.strip()}%"
        if search_by != "all" and search_by in ALL_MAAN_SERVICES_SEARCHABLE_COLS:
            where_clauses.append(f"LOWER({search_by}) LIKE LOWER(?)")
            params.append(q_like)
        else:
            or_clauses = [f"LOWER({c}) LIKE LOWER(?)" for c in ALL_MAAN_SERVICES_SEARCHABLE_COLS]
            where_clauses.append("(" + " OR ".join(or_clauses) + ")")
            params.extend([q_like] * len(ALL_MAAN_SERVICES_SEARCHABLE_COLS))
            
    if selected_type:
        where_clauses.append("LOWER(service_type) = LOWER(?)")
        params.append(selected_type.strip())
        
    where_sql = ""
    if where_clauses:
        where_sql = "WHERE " + " AND ".join(where_clauses)
        
    count_sql = f"SELECT COUNT(*) FROM maan_services {where_sql};"
    cursor.execute(count_sql, params)
    filtered_count = cursor.fetchone()[0]
    
    total_pages = max(1, (filtered_count + per_page - 1) // per_page)
    page = max(1, min(page, total_pages))
    offset = (page - 1) * per_page
    
    data_sql = f"SELECT * FROM maan_services {where_sql} ORDER BY {sort_col} {sort_dir} LIMIT ? OFFSET ?;"
    cursor.execute(data_sql, params + [per_page, offset])
    rows = [dict(r) for r in cursor.fetchall()]
    
    conn.close()
    
    return {
        'rows': rows,
        'total_records': total_records,
        'filtered_count': filtered_count,
        'unique_service_types': unique_service_types,
        'page': page,
        'per_page': per_page,
        'total_pages': total_pages,
        'query': query,
        'search_by': search_by,
        'selected_type': selected_type,
        'sort_by': sort_by,
        'sort_order': sort_order
    }


def search_maan_tls(query="", search_by="all", selected_media="", page=1, per_page=25, sort_by="id", sort_order="asc"):
    conn = get_db()
    cursor = conn.cursor()
    
    if search_by not in VALID_MAAN_TLS_SEARCH_COLUMNS:
        search_by = "all"
        
    sort_col = VALID_MAAN_TLS_SORT_COLUMNS.get(sort_by, 'maan_tls.id')
    sort_dir = "DESC" if str(sort_order).lower() in ("desc", "descending") else "ASC"
    
    cursor.execute("SELECT COUNT(*) FROM maan_tls;")
    total_records = cursor.fetchone()[0]
    
    cursor.execute("SELECT DISTINCT media_type FROM maan_tls WHERE media_type IS NOT NULL AND media_type != '' ORDER BY media_type ASC;")
    unique_media_types = [r['media_type'] for r in cursor.fetchall()]
    
    where_clauses = []
    params = []
    
    if query:
        q_like = f"%{query.strip()}%"
        if search_by != "all" and search_by in ALL_MAAN_TLS_SEARCHABLE_COLS:
            where_clauses.append(f"LOWER({search_by}) LIKE LOWER(?)")
            params.append(q_like)
        else:
            or_clauses = [f"LOWER({c}) LIKE LOWER(?)" for c in ALL_MAAN_TLS_SEARCHABLE_COLS]
            where_clauses.append("(" + " OR ".join(or_clauses) + ")")
            params.extend([q_like] * len(ALL_MAAN_TLS_SEARCHABLE_COLS))
            
    if selected_media:
        where_clauses.append("LOWER(media_type) = LOWER(?)")
        params.append(selected_media.strip())
        
    where_sql = ""
    if where_clauses:
        where_sql = "WHERE " + " AND ".join(where_clauses)
        
    count_sql = f"SELECT COUNT(*) FROM maan_tls {where_sql};"
    cursor.execute(count_sql, params)
    filtered_count = cursor.fetchone()[0]
    
    total_pages = max(1, (filtered_count + per_page - 1) // per_page)
    page = max(1, min(page, total_pages))
    offset = (page - 1) * per_page
    
    data_sql = f"SELECT * FROM maan_tls {where_sql} ORDER BY {sort_col} {sort_dir} LIMIT ? OFFSET ?;"
    cursor.execute(data_sql, params + [per_page, offset])
    rows = [dict(r) for r in cursor.fetchall()]
    
    conn.close()
    
    return {
        'rows': rows,
        'total_records': total_records,
        'filtered_count': filtered_count,
        'unique_media_types': unique_media_types,
        'page': page,
        'per_page': per_page,
        'total_pages': total_pages,
        'query': query,
        'search_by': search_by,
        'selected_media': selected_media,
        'sort_by': sort_by,
        'sort_order': sort_order
    }


# Page & API Routes for MAAN Nodes
@app.route('/maan-nodes')
@login_required
def maan_nodes():
    query = request.args.get('query', '').strip()
    search_by = request.args.get('search_by', 'all').strip()
    selected_ssa = request.args.get('ssa', '').strip()
    selected_type = request.args.get('type', '').strip()
    sort_by = request.args.get('sort_by', 'id').strip()
    sort_order = request.args.get('sort_order', 'asc').strip()
    try: page = int(request.args.get('page', 1))
    except ValueError: page = 1
    try:
        per_page = int(request.args.get('per_page', 25))
        if per_page not in (10, 25, 50, 100, 250, 500, 1000, 2000): per_page = 25
    except ValueError: per_page = 25
    res = search_maan_nodes(query, search_by, selected_ssa, selected_type, page, per_page, sort_by=sort_by, sort_order=sort_order)
    return render_template('maan_nodes.html', active_tab='nodes', nodes=res['rows'], total_records=res['total_records'], filtered_count=res['filtered_count'], unique_ssas=res['unique_ssas'], unique_types=res['unique_types'], unique_locations_count=res['unique_locations_count'], page=res['page'], per_page=res['per_page'], total_pages=res['total_pages'], query=query, search_by=search_by, selected_ssa=selected_ssa, selected_type=selected_type, sort_by=sort_by, sort_order=sort_order, search_columns=VALID_MAAN_SEARCH_COLUMNS)

@app.route('/api/maan-nodes/search')
@login_required
def api_search_maan_nodes():
    query = request.args.get('query', '').strip()
    search_by = request.args.get('search_by', 'all').strip()
    selected_ssa = request.args.get('ssa', '').strip()
    selected_type = request.args.get('type', '').strip()
    sort_by = request.args.get('sort_by', 'id').strip()
    sort_order = request.args.get('sort_order', 'asc').strip()
    try: page = int(request.args.get('page', 1))
    except ValueError: page = 1
    try:
        per_page = int(request.args.get('per_page', 25))
        if per_page not in (10, 25, 50, 100, 250, 500, 1000, 2000): per_page = 25
    except ValueError: per_page = 25
    res = search_maan_nodes(query, search_by, selected_ssa, selected_type, page, per_page, sort_by=sort_by, sort_order=sort_order)
    return jsonify(res)

@app.route('/api/maan-node/<int:node_id>')
@login_required
def get_maan_node(node_id):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM maan_nodes WHERE id = ?;", (node_id,))
    row = cursor.fetchone()
    conn.close()
    if row is None: return jsonify({'error': 'MAAN Node not found'}), 404
    return jsonify(dict(row))

@app.route('/maan-node/new', methods=['POST'])
@login_required
def create_maan_node():
    ne_ip = request.form.get('ne_ip', '').strip()
    location = request.form.get('location', '').strip()
    ne_type = request.form.get('type', '').strip()
    ssa = request.form.get('ssa', '').strip()
    phase = request.form.get('phase', '').strip()
    ne_name = request.form.get('ne_name', '').strip()
    dcc_ip = request.form.get('dcc_ip', '').strip()
    if not ne_ip or not location:
        flash('NE IP and Location are required fields.', 'warning')
        return redirect(url_for('maan_nodes'))
    if not ne_name: ne_name = f"{ne_ip}_{location}_{ne_type}_{ssa}".rstrip('_')
    conn = get_db()
    cursor = conn.cursor()
    try:
        cursor.execute('INSERT INTO maan_nodes (ne_ip, location, type, ssa, phase, ne_name, dcc_ip) VALUES (?, ?, ?, ?, ?, ?, ?);', (ne_ip, location, ne_type, ssa, phase, ne_name, dcc_ip))
        conn.commit()
        flash(f'MAAN Node {ne_ip} ({location}) added successfully!', 'success')
    except Exception as e:
        conn.rollback()
        flash(f'Error adding MAAN Node: {str(e)}', 'danger')
    finally: conn.close()
    return redirect(url_for('maan_nodes'))

@app.route('/maan-node/<int:node_id>/edit', methods=['POST'])
@login_required
def edit_maan_node(node_id):
    ne_ip = request.form.get('ne_ip', '').strip()
    location = request.form.get('location', '').strip()
    ne_type = request.form.get('type', '').strip()
    ssa = request.form.get('ssa', '').strip()
    phase = request.form.get('phase', '').strip()
    ne_name = request.form.get('ne_name', '').strip()
    dcc_ip = request.form.get('dcc_ip', '').strip()
    if not ne_ip or not location:
        flash('NE IP and Location are required fields.', 'warning')
        return redirect(url_for('maan_nodes'))
    conn = get_db()
    cursor = conn.cursor()
    try:
        cursor.execute('UPDATE maan_nodes SET ne_ip=?, location=?, type=?, ssa=?, phase=?, ne_name=?, dcc_ip=? WHERE id=?;', (ne_ip, location, ne_type, ssa, phase, ne_name, dcc_ip, node_id))
        conn.commit()
        flash(f'MAAN Node #{node_id} updated successfully!', 'success')
    except Exception as e:
        conn.rollback()
        flash(f'Error updating MAAN Node: {str(e)}', 'danger')
    finally: conn.close()
    return redirect(url_for('maan_nodes'))

@app.route('/maan-node/<int:node_id>/delete', methods=['POST'])
@login_required
def delete_maan_node(node_id):
    conn = get_db()
    cursor = conn.cursor()
    try:
        cursor.execute('DELETE FROM maan_nodes WHERE id=?;', (node_id,))
        conn.commit()
        flash(f'MAAN Node #{node_id} deleted successfully!', 'success')
    except Exception as e:
        conn.rollback()
        flash(f'Error deleting MAAN Node: {str(e)}', 'danger')
    finally: conn.close()
    return redirect(url_for('maan_nodes'))

@app.route('/upload-maan-nodes', methods=['POST'])
@login_required
def upload_maan_nodes():
    if 'file' not in request.files:
        flash('No file uploaded.', 'danger')
        return redirect(url_for('maan_nodes'))
    file = request.files['file']
    if file.filename == '':
        flash('No file selected.', 'warning')
        return redirect(url_for('maan_nodes'))
    try:
        if file.filename.endswith('.csv'): df = pd.read_csv(file)
        else: df = pd.read_excel(file)
        from init_db import import_maan_nodes_csv
        cnt = import_maan_nodes_csv(df)
        flash(f'Successfully imported {cnt} MAAN nodes!', 'success')
    except Exception as e: flash(f'Error processing file: {str(e)}', 'danger')
    return redirect(url_for('maan_nodes'))

@app.route('/export-maan-nodes')
@login_required
def export_maan_nodes():
    query = request.args.get('query', '').strip()
    search_by = request.args.get('search_by', 'all').strip()
    selected_ssa = request.args.get('ssa', '').strip()
    selected_type = request.args.get('type', '').strip()
    sort_by = request.args.get('sort_by', 'id').strip()
    sort_order = request.args.get('sort_order', 'asc').strip()
    export_format = request.args.get('format', 'csv').strip().lower()
    res = search_maan_nodes(query, search_by, selected_ssa, selected_type, page=1, per_page=20000, sort_by=sort_by, sort_order=sort_order)
    df = pd.DataFrame(res['rows'])
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output = BytesIO()
    if export_format in ('excel', 'xlsx'):
        filename = f"MAAN_Nodes_{timestamp}.xlsx"
        with pd.ExcelWriter(output, engine='openpyxl') as writer: df.to_excel(writer, index=False, sheet_name='MAAN Nodes')
        output.seek(0)
        return send_file(output, as_attachment=True, download_name=filename, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    else:
        filename = f"MAAN_Nodes_{timestamp}.csv"
        csv_bytes = df.to_csv(index=False, encoding='utf-8').encode('utf-8')
        output.write(csv_bytes)
        output.seek(0)
        return send_file(output, as_attachment=True, download_name=filename, mimetype='text/csv')


# Page & API Routes for MAAN Services
@app.route('/maan-services')
@login_required
def maan_services():
    query = request.args.get('query', '').strip()
    search_by = request.args.get('search_by', 'all').strip()
    selected_type = request.args.get('type', '').strip()
    sort_by = request.args.get('sort_by', 'id').strip()
    sort_order = request.args.get('sort_order', 'asc').strip()
    try: page = int(request.args.get('page', 1))
    except ValueError: page = 1
    try:
        per_page = int(request.args.get('per_page', 25))
        if per_page not in (10, 25, 50, 100, 250, 500, 1000, 2000): per_page = 25
    except ValueError: per_page = 25
    res = search_maan_services(query, search_by, selected_type, page, per_page, sort_by=sort_by, sort_order=sort_order)
    return render_template('maan_services.html', active_tab='services', rows=res['rows'], total_records=res['total_records'], filtered_count=res['filtered_count'], unique_service_types=res['unique_service_types'], page=res['page'], per_page=res['per_page'], total_pages=res['total_pages'], query=query, search_by=search_by, selected_type=selected_type, sort_by=sort_by, sort_order=sort_order, search_columns=VALID_MAAN_SERVICES_SEARCH_COLUMNS)

@app.route('/api/maan-services/search')
@login_required
def api_search_maan_services():
    query = request.args.get('query', '').strip()
    search_by = request.args.get('search_by', 'all').strip()
    selected_type = request.args.get('type', '').strip()
    sort_by = request.args.get('sort_by', 'id').strip()
    sort_order = request.args.get('sort_order', 'asc').strip()
    try: page = int(request.args.get('page', 1))
    except ValueError: page = 1
    try:
        per_page = int(request.args.get('per_page', 25))
        if per_page not in (10, 25, 50, 100, 250, 500, 1000, 2000): per_page = 25
    except ValueError: per_page = 25
    res = search_maan_services(query, search_by, selected_type, page, per_page, sort_by=sort_by, sort_order=sort_order)
    return jsonify(res)

@app.route('/api/maan-service/<int:service_id>')
@login_required
def get_maan_service(service_id):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM maan_services WHERE id = ?;", (service_id,))
    row = cursor.fetchone()
    conn.close()
    if row is None: return jsonify({'error': 'MAAN Service not found'}), 404
    return jsonify(dict(row))

@app.route('/maan-service/new', methods=['POST'])
@login_required
def create_maan_service():
    name = request.form.get('name', '').strip()
    service_type = request.form.get('service_type', '').strip()
    bandwidth_kbps = request.form.get('bandwidth_kbps', '0.0').strip()
    traffic_cir_kbps = request.form.get('traffic_cir_kbps', '0.0').strip()
    order_name = request.form.get('order_name', '').strip()
    client = request.form.get('client', '').strip()
    a_end = request.form.get('a_end', '').strip()
    z_end = request.form.get('z_end', '').strip()
    description = request.form.get('description', '').strip()
    if not name:
        flash('Service Name is required.', 'warning')
        return redirect(url_for('maan_services'))
    conn = get_db()
    cursor = conn.cursor()
    try:
        cursor.execute('INSERT INTO maan_services (name, service_type, bandwidth_kbps, traffic_cir_kbps, order_name, client, a_end, z_end, description) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?);', (name, service_type, bandwidth_kbps, traffic_cir_kbps, order_name, client, a_end, z_end, description))
        conn.commit()
        flash(f'MAAN Service "{name}" added successfully!', 'success')
    except Exception as e:
        conn.rollback()
        flash(f'Error adding MAAN Service: {str(e)}', 'danger')
    finally: conn.close()
    return redirect(url_for('maan_services'))

@app.route('/maan-service/<int:service_id>/edit', methods=['POST'])
@login_required
def edit_maan_service(service_id):
    name = request.form.get('name', '').strip()
    service_type = request.form.get('service_type', '').strip()
    bandwidth_kbps = request.form.get('bandwidth_kbps', '0.0').strip()
    traffic_cir_kbps = request.form.get('traffic_cir_kbps', '0.0').strip()
    order_name = request.form.get('order_name', '').strip()
    client = request.form.get('client', '').strip()
    a_end = request.form.get('a_end', '').strip()
    z_end = request.form.get('z_end', '').strip()
    description = request.form.get('description', '').strip()
    if not name:
        flash('Service Name is required.', 'warning')
        return redirect(url_for('maan_services'))
    conn = get_db()
    cursor = conn.cursor()
    try:
        cursor.execute('UPDATE maan_services SET name=?, service_type=?, bandwidth_kbps=?, traffic_cir_kbps=?, order_name=?, client=?, a_end=?, z_end=?, description=? WHERE id=?;', (name, service_type, bandwidth_kbps, traffic_cir_kbps, order_name, client, a_end, z_end, description, service_id))
        conn.commit()
        flash(f'MAAN Service #{service_id} updated successfully!', 'success')
    except Exception as e:
        conn.rollback()
        flash(f'Error updating MAAN Service: {str(e)}', 'danger')
    finally: conn.close()
    return redirect(url_for('maan_services'))

@app.route('/maan-service/<int:service_id>/delete', methods=['POST'])
@login_required
def delete_maan_service(service_id):
    conn = get_db()
    cursor = conn.cursor()
    try:
        cursor.execute('DELETE FROM maan_services WHERE id=?;', (service_id,))
        conn.commit()
        flash(f'MAAN Service #{service_id} deleted successfully!', 'success')
    except Exception as e:
        conn.rollback()
        flash(f'Error deleting MAAN Service: {str(e)}', 'danger')
    finally: conn.close()
    return redirect(url_for('maan_services'))

@app.route('/upload-maan-services', methods=['POST'])
@login_required
def upload_maan_services():
    if 'file' not in request.files:
        flash('No file uploaded.', 'danger')
        return redirect(url_for('maan_services'))
    file = request.files['file']
    if file.filename == '':
        flash('No file selected.', 'warning')
        return redirect(url_for('maan_services'))
    try:
        if file.filename.endswith('.csv'): df = pd.read_csv(file)
        else: df = pd.read_excel(file)
        from init_db import import_maan_services_csv
        cnt = import_maan_services_csv(df)
        flash(f'Successfully imported {cnt} MAAN services!', 'success')
    except Exception as e: flash(f'Error processing file: {str(e)}', 'danger')
    return redirect(url_for('maan_services'))

@app.route('/export-maan-services')
@login_required
def export_maan_services():
    query = request.args.get('query', '').strip()
    search_by = request.args.get('search_by', 'all').strip()
    selected_type = request.args.get('type', '').strip()
    sort_by = request.args.get('sort_by', 'id').strip()
    sort_order = request.args.get('sort_order', 'asc').strip()
    export_format = request.args.get('format', 'csv').strip().lower()
    res = search_maan_services(query, search_by, selected_type, page=1, per_page=20000, sort_by=sort_by, sort_order=sort_order)
    df = pd.DataFrame(res['rows'])
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output = BytesIO()
    if export_format in ('excel', 'xlsx'):
        filename = f"MAAN_Services_{timestamp}.xlsx"
        with pd.ExcelWriter(output, engine='openpyxl') as writer: df.to_excel(writer, index=False, sheet_name='MAAN Services')
        output.seek(0)
        return send_file(output, as_attachment=True, download_name=filename, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    else:
        filename = f"MAAN_Services_{timestamp}.csv"
        csv_bytes = df.to_csv(index=False, encoding='utf-8').encode('utf-8')
        output.write(csv_bytes)
        output.seek(0)
        return send_file(output, as_attachment=True, download_name=filename, mimetype='text/csv')


# Page & API Routes for MAAN TLS
@app.route('/maan-tls')
@login_required
def maan_tls():
    query = request.args.get('query', '').strip()
    search_by = request.args.get('search_by', 'all').strip()
    selected_media = request.args.get('media', '').strip()
    sort_by = request.args.get('sort_by', 'id').strip()
    sort_order = request.args.get('sort_order', 'asc').strip()
    try: page = int(request.args.get('page', 1))
    except ValueError: page = 1
    try:
        per_page = int(request.args.get('per_page', 25))
        if per_page not in (10, 25, 50, 100, 250, 500, 1000, 2000): per_page = 25
    except ValueError: per_page = 25
    res = search_maan_tls(query, search_by, selected_media, page, per_page, sort_by=sort_by, sort_order=sort_order)
    return render_template('maan_tls.html', active_tab='tls', rows=res['rows'], total_records=res['total_records'], filtered_count=res['filtered_count'], unique_media_types=res['unique_media_types'], page=res['page'], per_page=res['per_page'], total_pages=res['total_pages'], query=query, search_by=search_by, selected_media=selected_media, sort_by=sort_by, sort_order=sort_order, search_columns=VALID_MAAN_TLS_SEARCH_COLUMNS)

@app.route('/api/maan-tls/search')
@login_required
def api_search_maan_tls():
    query = request.args.get('query', '').strip()
    search_by = request.args.get('search_by', 'all').strip()
    selected_media = request.args.get('media', '').strip()
    sort_by = request.args.get('sort_by', 'id').strip()
    sort_order = request.args.get('sort_order', 'asc').strip()
    try: page = int(request.args.get('page', 1))
    except ValueError: page = 1
    try:
        per_page = int(request.args.get('per_page', 25))
        if per_page not in (10, 25, 50, 100, 250, 500, 1000, 2000): per_page = 25
    except ValueError: per_page = 25
    res = search_maan_tls(query, search_by, selected_media, page, per_page, sort_by=sort_by, sort_order=sort_order)
    return jsonify(res)

@app.route('/api/maan-tls/<int:tls_id>')
@login_required
def get_maan_tls(tls_id):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM maan_tls WHERE id = ?;", (tls_id,))
    row = cursor.fetchone()
    conn.close()
    if row is None: return jsonify({'error': 'MAAN TLS Record not found'}), 404
    return jsonify(dict(row))

@app.route('/maan-tls/new', methods=['POST'])
@login_required
def create_maan_tls():
    name = request.form.get('name', '').strip()
    media_type = request.form.get('media_type', '').strip()
    bandwidth = request.form.get('bandwidth', '').strip()
    signal_type = request.form.get('signal_type', '').strip()
    a_end = request.form.get('a_end', '').strip()
    z_end = request.form.get('z_end', '').strip()
    alarm_status = request.form.get('alarm_status', 'Normal').strip()
    order_name = request.form.get('order_name', '').strip()
    client = request.form.get('client', '').strip()
    description = request.form.get('description', '').strip()
    if not name:
        flash('Name is required.', 'warning')
        return redirect(url_for('maan_tls'))
    conn = get_db()
    cursor = conn.cursor()
    try:
        cursor.execute('INSERT INTO maan_tls (name, media_type, bandwidth, signal_type, a_end, z_end, alarm_status, order_name, client, description) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?);', (name, media_type, bandwidth, signal_type, a_end, z_end, alarm_status, order_name, client, description))
        conn.commit()
        flash(f'MAAN TLS Record "{name}" added successfully!', 'success')
    except Exception as e:
        conn.rollback()
        flash(f'Error adding MAAN TLS Record: {str(e)}', 'danger')
    finally: conn.close()
    return redirect(url_for('maan_tls'))

@app.route('/maan-tls/<int:tls_id>/edit', methods=['POST'])
@login_required
def edit_maan_tls(tls_id):
    name = request.form.get('name', '').strip()
    media_type = request.form.get('media_type', '').strip()
    bandwidth = request.form.get('bandwidth', '').strip()
    signal_type = request.form.get('signal_type', '').strip()
    a_end = request.form.get('a_end', '').strip()
    z_end = request.form.get('z_end', '').strip()
    alarm_status = request.form.get('alarm_status', 'Normal').strip()
    order_name = request.form.get('order_name', '').strip()
    client = request.form.get('client', '').strip()
    description = request.form.get('description', '').strip()
    if not name:
        flash('Name is required.', 'warning')
        return redirect(url_for('maan_tls'))
    conn = get_db()
    cursor = conn.cursor()
    try:
        cursor.execute('UPDATE maan_tls SET name=?, media_type=?, bandwidth=?, signal_type=?, a_end=?, z_end=?, alarm_status=?, order_name=?, client=?, description=? WHERE id=?;', (name, media_type, bandwidth, signal_type, a_end, z_end, alarm_status, order_name, client, description, tls_id))
        conn.commit()
        flash(f'MAAN TLS Record #{tls_id} updated successfully!', 'success')
    except Exception as e:
        conn.rollback()
        flash(f'Error updating MAAN TLS Record: {str(e)}', 'danger')
    finally: conn.close()
    return redirect(url_for('maan_tls'))

@app.route('/maan-tls/<int:tls_id>/delete', methods=['POST'])
@login_required
def delete_maan_tls(tls_id):
    conn = get_db()
    cursor = conn.cursor()
    try:
        cursor.execute('DELETE FROM maan_tls WHERE id=?;', (tls_id,))
        conn.commit()
        flash(f'MAAN TLS Record #{tls_id} deleted successfully!', 'success')
    except Exception as e:
        conn.rollback()
        flash(f'Error deleting MAAN TLS Record: {str(e)}', 'danger')
    finally: conn.close()
    return redirect(url_for('maan_tls'))

@app.route('/upload-maan-tls', methods=['POST'])
@login_required
def upload_maan_tls():
    if 'file' not in request.files:
        flash('No file uploaded.', 'danger')
        return redirect(url_for('maan_tls'))
    file = request.files['file']
    if file.filename == '':
        flash('No file selected.', 'warning')
        return redirect(url_for('maan_tls'))
    try:
        if file.filename.endswith('.csv'): df = pd.read_csv(file)
        else: df = pd.read_excel(file)
        from init_db import import_maan_tls_csv
        cnt = import_maan_tls_csv(df)
        flash(f'Successfully imported {cnt} MAAN TLS records!', 'success')
    except Exception as e: flash(f'Error processing file: {str(e)}', 'danger')
    return redirect(url_for('maan_tls'))

@app.route('/export-maan-tls')
@login_required
def export_maan_tls():
    query = request.args.get('query', '').strip()
    search_by = request.args.get('search_by', 'all').strip()
    selected_media = request.args.get('media', '').strip()
    sort_by = request.args.get('sort_by', 'id').strip()
    sort_order = request.args.get('sort_order', 'asc').strip()
    export_format = request.args.get('format', 'csv').strip().lower()
    res = search_maan_tls(query, search_by, selected_media, page=1, per_page=20000, sort_by=sort_by, sort_order=sort_order)
    df = pd.DataFrame(res['rows'])
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output = BytesIO()
    if export_format in ('excel', 'xlsx'):
        filename = f"MAAN_TLS_{timestamp}.xlsx"
        with pd.ExcelWriter(output, engine='openpyxl') as writer: df.to_excel(writer, index=False, sheet_name='MAAN TLS')
        output.seek(0)
        return send_file(output, as_attachment=True, download_name=filename, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    else:
        filename = f"MAAN_TLS_{timestamp}.csv"
        csv_bytes = df.to_csv(index=False, encoding='utf-8').encode('utf-8')
        output.write(csv_bytes)
        output.seek(0)
        return send_file(output, as_attachment=True, download_name=filename, mimetype='text/csv')


# ==============================================================================

if __name__ == '__main__':
    init_db()
    print("Starting CPAN Network on http://127.0.0.1:5000 ...")
    app.run(host='127.0.0.1', port=5000, debug=True)




