import os
import gc
import openpyxl
import pandas as pd

bts_path = r'c:\Users\Admin\Desktop\txsystem\btsdatabase.xlsx'
master_path = r'c:\Users\Admin\Desktop\txsystem\Master_Record_TCS.xlsx'
temp_path = r'c:\Users\Admin\Desktop\txsystem\btsdatabase_temp.xlsx'

print("Loading Master_Record_TCS.xlsx...")
df_master = pd.read_excel(master_path)

# Build lookup maps
lookup_by_bts_id = {}
lookup_by_bts_name = {}

for idx, row in df_master.iterrows():
    bts_id = str(row['BTS ID']).strip().upper() if pd.notna(row['BTS ID']) else ""
    bts_name = str(row['BTS Name']).strip().upper() if pd.notna(row['BTS Name']) else ""
    
    data = {
        'tx-system-ip': str(row['Endpoint Node IP']).strip() if pd.notna(row['Endpoint Node IP']) and str(row['Endpoint Node IP']).strip().lower() != 'nan' else "",
        'tx-system-location': str(row['Endpoint Node']).strip() if pd.notna(row['Endpoint Node']) and str(row['Endpoint Node']).strip().lower() != 'nan' else "",
        'tx-system-port': str(row['Endpoint Port']).strip() if pd.notna(row['Endpoint Port']) and str(row['Endpoint Port']).strip().lower() != 'nan' else ""
    }
    
    if bts_id and bts_id != 'NAN':
        lookup_by_bts_id[bts_id] = data
    if bts_name and bts_name != 'NAN':
        lookup_by_bts_name[bts_name] = data

print(f"Loaded {len(lookup_by_bts_id)} BTS IDs from Master_Record_TCS.xlsx")

# Load btsdatabase.xlsx
with open(bts_path, 'rb') as f:
    wb = openpyxl.load_workbook(f)

sheet = wb['Sheet1']

headers = [sheet.cell(row=1, column=c).value for c in range(1, sheet.max_column + 1)]
site_id_idx = headers.index('site_id') + 1
site_name_idx = headers.index('site_name') + 1
ssa_idx = headers.index('ssa') + 1

tx_ip_idx = headers.index('tx-system-ip') + 1
tx_loc_idx = headers.index('tx-system-location') + 1
tx_port_idx = headers.index('tx-system-port') + 1

updated_count = 0

for r in range(2, sheet.max_row + 1):
    site_id_val = str(sheet.cell(row=r, column=site_id_idx).value or '').strip().upper()
    site_name_val = str(sheet.cell(row=r, column=site_name_idx).value or '').strip().upper()
    ssa_val = str(sheet.cell(row=r, column=ssa_idx).value or '').strip().upper()
    
    data = None
    if site_id_val in lookup_by_bts_id:
        data = lookup_by_bts_id[site_id_val]
    elif '_' in site_name_val and site_name_val.split('_', 1)[0] in lookup_by_bts_id:
        data = lookup_by_bts_id[site_name_val.split('_', 1)[0]]
    elif site_name_val in lookup_by_bts_name:
        data = lookup_by_bts_name[site_name_val]
    elif ssa_val in lookup_by_bts_id:
        data = lookup_by_bts_id[ssa_val]
    else:
        # Match substring BTS ID inside site_id
        for b_id, b_data in lookup_by_bts_id.items():
            if b_id and (site_id_val.endswith(b_id) or site_name_val.startswith(b_id)):
                data = b_data
                break
                
    if data:
        sheet.cell(row=r, column=tx_ip_idx, value=data['tx-system-ip'])
        sheet.cell(row=r, column=tx_loc_idx, value=data['tx-system-location'])
        sheet.cell(row=r, column=tx_port_idx, value=data['tx-system-port'])
        updated_count += 1

wb.save(temp_path)
wb.close()
del wb
gc.collect()

os.replace(temp_path, bts_path)
print(f"Successfully updated VLOOKUP transmission data for {updated_count} / {sheet.max_row - 1} records in btsdatabase.xlsx!")

# Re-import into SQLite relational database
from init_db import init_db
init_db(force_reimport=True)
print("Re-imported updated btsdatabase.xlsx into SQLite database (btsdatabase.db) successfully!")
